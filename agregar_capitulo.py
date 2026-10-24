# -*- coding: utf-8 -*-
"""
===============================================================================
GESTOR UNIFICADO DE CAPÍTULOS - "EL CAMINO DEL REENCARNADO"
===============================================================================
Este archivo es la plantilla GENERAL y ÚNICA para agregar o sobrescribir capítulos.
Para agregar un nuevo capítulo, simplemente edite el diccionario `CAPITULO` con
el número, título y párrafos correspondientes y ejecute este script:
    python agregar_capitulo.py

El script automáticamente:
1. Genera el documento Word individual: capitulos/XX_Capitulo_XX.docx
2. Recompila el libro maestro completo: compilado/El_Camino_del_Reencarnado.docx
===============================================================================
"""

import os
import sys
import glob
import re
import docx

# Definición de rutas base
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
PROYECTO_DIR = r"c:\Users\fsilv\OneDrive\Documentos\GitHub\Proyecto-Minecraft\Novela"

for d in [CURRENT_DIR, NOVELA_DIR, PROYECTO_DIR]:
    if os.path.exists(d) and d not in sys.path:
        sys.path.insert(0, d)

from generar_libro import create_novel_document, create_single_chapter_document

# ===============================================================================
# DEFINICIÓN DEL CAPÍTULO A AGREGAR / ACTUALIZAR
# ===============================================================================
CAPITULO = {   'num': 87,
    'header': 'Capítulo 87',
    'title': 'El Enigma del Portón Imperial: El Segundo Círculo y la Cripta de Disformidad',
    'paragraphs': [   'El silencio que siguió a la calcinación del primer saco embrionario era denso y cargado de '
                      'expectación.',
                      'Sobre el lecho del antiguo acueducto de Valdemar, las aguas purificadas corrían con un murmullo '
                      'fresco y transparente, pero a escasos veinte metros de nosotros se alzaba el verdadero desafío '
                      'de la expedición: el gigantesco portón ciclópeo de piedra negra y basalto, incrustado en la '
                      'roca viva con la solidez de una montaña.',
                      'Los relieves ancestrales que adornaban el dintel representaban las insignias heráldicas de la '
                      'Tercera Dinastía de Solaria —soles alados, leones rampantes y espadas cruzadas—, pero la '
                      'superficie estaba mancillada por gruesas costras de miasma petrificado y pátinas de un verde '
                      'esmeralda ponzoñoso que chisporroteaban con un tenue resplandor fosforescente.',
                      '—Es una cerradura cuádruple de engranajes contrarrotatorios con traba rúnica de sellado '
                      'imperial —dictaminó el Capitán Bawden, arrodillándose frente al mecanismo central mientras '
                      'pasaba la hoja de su daga de plata por las ranuras del metal—. En los archivos de la Guardia '
                      'Plateada estudiamos estos portones. Si se fuerza el cilindro exterior sin desalinear los '
                      'pestillos magnéticos, las runas defensivas detonarán un colapso de caverna que nos sepultará a '
                      'todos.',
                      '—Y la corrosión por piedra bruja ha soldado los pernos de mithril —añadió Agnes Tachyon, '
                      'ajustándose sus lentes de alquimista con una sonrisa de emoción científica—. Pero nada que una '
                      'solución endotérmica de ácido fluorhídrico y extracto de loto no pueda disolver en sesenta '
                      'segundos.',
                      'Agnes extrajo de su bandolera un frasco de cristal templado y roció un líquido azul pálido '
                      'sobre las ranuras oxidadas. El miasma petrificado silbó con una efervescencia gélida, '
                      'disolviéndose en una humareda blanca inodora que dejó los engranajes relucientes.',
                      'Bawden alineó los tres primeros tambores mecánicos con movimientos certeros de su estilete. '
                      'Solo faltaba el giro maestro.',
                      'Di un paso al frente y aferré el manubrio de hierro macizo. Canalicé maná telúrico a través de '
                      'mis antebrazos: los brazaletes forjados en Morvath vibraron con un tono grave y resonante, '
                      'transmitiendo una fuerza bruta controlada hacia los ejes internos. Con un crujido monumental de '
                      'goznes milenarios que sacudió el polvo de las bóvedas, el portón ciclópeo se desplazó hacia el '
                      'interior de la roca, abriéndose de par en par.',
                      '* * *',
                      'Una bocanada de aire abrasador nos golpeó de lleno en el rostro al cruzar el umbral.',
                      'Habíamos descendido al **Segundo Círculo Subterráneo** de las Ruinas Negras. Aquí, la frialdad '
                      'de las galerías de ónix había desaparecido por completo: la temperatura ascendía drásticamente '
                      'y la atmósfera olía a azufre, ozono electrificado y el acre hedor metálico de la piedra bruja '
                      'concentrada.',
                      'A ambos lados del camino, estrechas canaletas talladas en el suelo transportaban hilos de '
                      'líquido verde luminiscente que palpitaban como venas expuestas de un organismo aberrante. El '
                      'visor de mi respirador de loto se empañó levemente por el calor, pero el filtro de Lortelle y '
                      'el broche gélido de Lady Elysea en mi pecho mantuvieron mi respiración fresca y mi cuerpo '
                      'blindado contra la radiación.',
                      '—¡P-Por todos los dioses marciales...! —gimió Kenichi Shirahama a mi espalda, tapándose la boca '
                      'con ambas manos enguantadas—. ¡Ese charco verde brilla exactamente igual que la sopa de algas '
                      "fermentadas que el maestro Kensei Ma me obligaba a tragar para 'endurecer las tripas'! ¡Siento "
                      'que si doy un paso en falso me saldrán tres ojos!',
                      'Subaru Natsuki soltó una carcajada apagada detrás de su máscara:',
                      '—No la bebas, Kenichi, a menos que quieras que tu *Seikūken* dispare rayos por las orejas. '
                      'Mantente pegado a la retaguardia de Erza y vigila tus esquinas.',
                      '—Silencio en la columna —ordenó Erza con severidad marcial, aunque una sombra de sonrisa se '
                      'adivinó bajo su mirada—. Entramos en sector hostil fortificado.',
                      'Al avanzar cincuenta metros más por el desfiladero iluminado por el miasma verde, la galería se '
                      'ensanchó abruptamente hacia una nave colosal: una antigua cripta imperial de techos abovedados '
                      'y columnas corintias mutiladas, reconvertida en una siniestra fundición de guerra Skaven.',
                      'Habíamos localizado la **Segunda Cámara Nodriza**.',
                      '* * *',
                      'En el centro de la cripta, suspendido a diez metros de altura sobre un foso de piedra bruja, '
                      'latía un segundo Saco Embrionario Nodriza, mucho más grande y grotesco que el anterior. '
                      'Mangueras de cobre y tuberías de cristal grueso conectaban la masa cáustica con tres '
                      'generadores de disformidad cilíndricos que zumbaban con chispazos eléctricos esmeralda.',
                      'Pero la sala no estaba desprotegida.',
                      'Alrededor del foso se desplegaba una guarnición de élite del Clan Skryre: más de cincuenta '
                      '**Alimañas Acorazadas** con cotas de malla de hierro negro y picas electrificadas, respaldadas '
                      'en las galerías superiores por tiradores con mosquetes de chispa Jezzails y dos parejas de '
                      'ingenieros brujos operando lanzallamas de disformidad con mochilas presurizadas.',
                      'En cuanto nuestros pasos hicieron crujir el basalto, un chirrido histérico y gutural resonó en '
                      'la bóveda:',
                      '—¡In-intrusos-carne de Solaria! —bramó el Ingeniero Brujo Mayor, un roedor encorvado con un ojo '
                      'biónico de bronce y un cetro electrificado—. ¡Carguen-disparen los tubos de chispa! ¡Háganlos '
                      'ceniza-chatarra!',
                      'Los cañones de los mosquetes Jezzails relampaguearon en las cornisas con un estruendo de '
                      'truenos esmeralda.',
                      '—¡Bastión de Hierro, formación tortuga! —rugió el Capitán Kaelen Thorne.',
                      'Los soldados de tercer año plantaron sus gigantescos paveses de granito reforzado con mithril '
                      'en una muralla impenetrable. Los proyectiles de piedra bruja impactaron contra los escudos con '
                      'explosiones de chispas verdes, pero las runas de tierra de Thorne dispersaron la carga '
                      'eléctrica directamente hacia el suelo sin que un solo cadete retrocediera un milímetro.',
                      '—¡Dante, Hoshina, limpien las alturas! —ordenó Seier.',
                      'Dante Vance apoyó su rifle rúnico sobre el hombro de uno de los escuderos de Thorne. Con una '
                      'frialdad absoluta, ajustó la mira óptica y apretó el gatillo tres veces seguidas: los '
                      'estampidos rúnicos atravesaron la penumbra con destellos dorados, haciendo caer a los tres '
                      'francotiradores Skaven desde las cornisas superiores con impactos limpios en el pecho.',
                      'A su lado, Subaru Hoshina disparó en ráfagas de dos proyectiles perforantes, inutilizando los '
                      'mecanismos de disparo de los tiradores restantes y obligando a los hostigadores a replegarse '
                      'tras las columnas.',
                      '* * *',
                      '—¡Avanzan las picas! —gritó Erza mientras invocaba seis espadas flotantes a su alrededor—. '
                      '¡Lighter, Kenichi, desarticulen la primera línea!',
                      'Las Alimañas Acorazadas cargaron en falange cerrada, bajando sus picas electrificadas para '
                      'empalarnos.',
                      'Lighter no dudó un segundo: flexionó las rodillas, encendió los propulsores térmicos de sus '
                      'guanteletes y se lanzó en carrera recta como un cometa incandescente. Con una serie de jabs '
                      'supersónicos de combustión directa (*Flash Burst*), Lighter astilló las picas centrales y '
                      'conectó ganchos fulminantes a la mandíbula de tres guardias, haciéndolos volar hacia atrás '
                      'envueltos en llamaradas de vapor.',
                      'Por el flanco izquierdo, la falange enemiga intentó cerrar el hueco lanzando tres lanzas '
                      'directamente hacia la cabeza de Kenichi.',
                      'Lo que ocurrió a continuación fue una lección magistral de artes marciales orientales.',
                      'Kenichi Shirahama respiró hondo y expandió su ***Seikūken***. Sus ojos castaños leyeron las '
                      'trayectorias de las puntas de acero antes de que alcanzaran su espacio vital. Sin retroceder un '
                      'milímetro, ladeó el torso con la fluidez del agua, desvió la primera lanza con el canto '
                      'exterior de su antebrazo, atrapó el asta de la segunda con un giro milimétrico de muñeca (*Kote '
                      'Gaeshi*) y utilizó la inercia de la carga enemiga para ejecutar un *Ippon Seoi Nage* colosal.',
                      'El soldado rata acorazado de casi cien kilos sobrevoló su hombro describiendo una parábola '
                      'perfecta y se estrelló de lleno contra el equipo del lanzallamas de disformidad, rompiendo la '
                      'manguera presurizada en una fuga inofensiva de gas verde.',
                      'Sin detener su compás, Kenichi giró sobre su talón, barrió las piernas de otros dos piqueros '
                      'con una patada baja de barrido y los proyectó contra los adoquines con una palanca articular '
                      'que los desarmó al instante.',
                      '—¡Ufff, qué miedo! ¡Esas lanzas tenían corriente eléctrica de verdad! —gritó Kenichi secándose '
                      'la frente con el dorso del guante, mientras dejaba a cuatro alimañas gimiendo en el suelo sin '
                      'haber matado a ninguna.',
                      'Lighter le dedicó un guiño cómplice y un pulgar arriba entre el humo:',
                      '—¡Hermoso juego de cadera, muchacho! ¡Tienes más cintura que un boxeador de campeonato!',
                      '* * *',
                      '—¡Seier, el generador maestro va a sobrecargarse! —gritó Agnes Tachyon señalando al Ingeniero '
                      'Brujo Mayor, que manipulaba frenéticamente las manivelas del cañón central para desatar un rayo '
                      'masivo hacia la nave.',
                      '—¡Vamos, Ken! —rugí, envolviendo mis piernas y puños en una densa coraza de maná telúrico de '
                      'titanio.',
                      'Arremetí hacia el estrado central esquivando esquirlas y chispas de piedra bruja. El roedor '
                      'líder intentó apuntarme con su cetro electrificado, pero acorté la distancia con un paso '
                      'explosivo de Jeet Kune Do y descargué un codazo horizontal (*Sok Ti*) potenciado por el hierro '
                      'de Morvath directo a la mochila generadora de su espalda.',
                      'El impacto reventó los engranajes de cobre y los cables mágicos en una detonación de humo negro '
                      'que arrojó al ingeniero contra las columnas de la cripta, neutralizándolo por completo.',
                      'En ese mismo instante, una sombra surcó el aire por encima de mi cabeza.',
                      'Ken se impulsó desde mis hombros con una fuerza formidable, ascendiendo hacia la bóveda con su '
                      'bandana roja flameando entre chispas doradas. En el cenit de su salto, concentró toda su '
                      'energía marcial en su pierna derecha, envolviéndola en una estela de fuego solar puro y '
                      'ardiente como el mediodía:',
                      '—¡ARDE EN LAS CENIZAS DEL FÉNIX! —bramó el artista marcial con voz desgarradora.',
                      'Descargó una patada de hacha solar descendente (*Ignición Solar Purificadora*) que se clavó '
                      'como un cometa en el corazón del segundo Saco Embrionario Nodriza. La deflagración solar fue '
                      'cegadora: una explosión blanca y dorada de calor estelar puro que desintegró la biomasa '
                      'purpúrea, vaporizó las venas de miasma y fundió los depósitos de piedra bruja a nivel molecular '
                      'sin permitir que escapara una sola espora tóxica.',
                      'La cripta entera tembló mientras la podredumbre se convertía en una nube de ceniza inerte que '
                      'cayó sobre el foso como nieve apagada.',
                      '* * *',
                      'El zumbido de los generadores murió. El miasma ambiental se desplomó drásticamente, permitiendo '
                      'que la luz de nuestras linteras rúnicas iluminara con nitidez la majestuosidad de la '
                      'arquitectura imperial antigua.',
                      'Agnes Tachyon corrió hacia los restos del generador con pinzas de aleación y tubos de ensayo:',
                      '—¡La toxicidad ambiental cayó otro treinta y cinco por ciento! ¡Estamos en casi un ochenta por '
                      'ciento de purificación total de las ruinas! ¡Y miren estos cristales de disformidad purificados '
                      'por el fuego solar, sus matrices atómicas son absolutamente fascinantes!',
                      'El Capitán Bawden se acercó con su capa ondeando, envainó su estoque y me miró con una sonrisa '
                      'cargada de camaradería y respeto genuino:',
                      '—Trabajo impecable, Silva. Desarmar una falange electrificada y un generador Skryre en menos de '
                      'cinco minutos sin sufrir una sola baja... Gerald no exageraba cuando dijo que su generación '
                      'estaba forjada en otro molde.',
                      '—Y no olviden al chico de Hinomoto —añadió Thorne con una carcajada bonachona, palmeando la '
                      'espalda de Kenichi hasta hacerlo trastabillar con su petate—. Ese truco de voltear al piquero '
                      'sobre el lanzallamas nos ahorró diez bajas seguras. Bien hecho, Shirahama.',
                      'Kenichi se rascó la cabeza colorado como un tomate, mientras Subaru y Dante se reían con él.',
                      'Me acerqué al fondo de la gran nave. Allí, donde la piedra bruja había sido calcinada, el suelo '
                      'cedía en una rampa helicoidal de peldaños de granito que descendía hacia la oscuridad más '
                      'profunda del complejo: el **Tercer Círculo Subterráneo**, el corazón final de las Ruinas Negras '
                      'donde los mapas de Lortelle situaban la Cámara del Soberano de Miasma y la fuente raíz de la '
                      'corrupción.',
                      'Apreté el broche de Lady Elysea y el amuleto de Kaoruko en mi pecho, sintiendo el peso de la '
                      'misión y la calidez de su recuerdo.',
                      '—Descansen tres minutos y recarguen munición —dije alzando la voz con firmeza y calma ante los '
                      'Cuatro Estandartes—. La segunda nodriza está destruida. El corazón del nido nos espera abajo.']}


# ===============================================================================
# FUNCIONES DE GESTIÓN Y COMPILACIÓN
# ===============================================================================

def guardar_capitulo_individual(cap_dict, base_dir=NOVELA_DIR):
    """Guarda el archivo .docx individual del capítulo."""
    capitulos_dir = os.path.join(base_dir, "capitulos")
    os.makedirs(capitulos_dir, exist_ok=True)
    
    num = cap_dict.get("num", 0)
    if num == 0 or "PROLOGO" in cap_dict.get("header", "").upper() or "PRÓLOGO" in cap_dict.get("header", "").upper():
        base_name = "00_Prologo"
    else:
        base_name = f"{num:02d}_Capitulo_{num}"
        
    docx_path = os.path.join(capitulos_dir, f"{base_name}.docx")
    create_single_chapter_document(
        cap_dict,
        docx_path,
        novel_title="El camino del reencarnado",
        author="Seier Silva"
    )
    print(f"[OK] Documento Word individual guardado:\n  {docx_path}")
    return docx_path


def cargar_todos_los_capitulos(base_dir=NOVELA_DIR):
    """Lee todos los capítulos directamente desde la carpeta capitulos/*.docx."""
    capitulos_dir = os.path.join(base_dir, "capitulos")
    docx_files = sorted(glob.glob(os.path.join(capitulos_dir, "*.docx")))
    
    # Intentar cargar mapa de ilustraciones
    ilustraciones = {}
    try:
        from ilustraciones_capitulos import ILUSTRACIONES_POR_CAPITULO
        ilustraciones = ILUSTRACIONES_POR_CAPITULO
    except Exception:
        pass
        
    lista_capitulos = []
    
    for fpath in docx_files:
        filename = os.path.basename(fpath)
        num = None
        m_num = re.match(r"^(\d+)_", filename)
        if m_num:
            num = int(m_num.group(1))
            
        doc = docx.Document(fpath)
        if not doc.paragraphs:
            continue
            
        title_p = doc.paragraphs[0].text.strip()
        parts = title_p.split("\n")
        header = parts[0].strip() if len(parts) > 0 else "Capítulo"
        title = parts[1].strip() if len(parts) > 1 else ""
        
        paragraphs = []
        for p in doc.paragraphs[1:]:
            t = p.text.strip()
            if not t:
                continue
            if t in ["— ✦ —", "— ◆ —", "✦ ✦ ✦", "— ✧ —"]:
                continue
            if t.startswith("Ilustración:") or t.startswith("Ilustracion:"):
                continue
            paragraphs.append(t)
            
        chap_dict = {
            "header": header,
            "title": title,
            "paragraphs": paragraphs
        }
        
        # Asociar ilustraciones
        if num in ilustraciones:
            illus_info = ilustraciones[num]
            if "illustrations" in illus_info:
                chap_dict["illustrations"] = illus_info["illustrations"]
            elif "illustration" in illus_info:
                chap_dict["illustration"] = illus_info["illustration"]
                chap_dict["illustration_caption"] = illus_info.get("illustration_caption", "")
                
        lista_capitulos.append(chap_dict)
        
    return lista_capitulos


def compilar_libro_completo(base_dir=NOVELA_DIR):
    """Recompila el libro maestro completo a partir de todos los .docx de capitulos/."""
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    
    capitulos = cargar_todos_los_capitulos(base_dir)
    print(f"[*] Compilando Libro Maestro con {len(capitulos)} capítulos (desde .docx)...")
    
    compilado_dir = os.path.join(base_dir, "compilado")
    os.makedirs(compilado_dir, exist_ok=True)
    master_docx = os.path.join(compilado_dir, "El_Camino_del_Reencarnado.docx")
    
    try:
        create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
        print(f"[OK] Libro Maestro actualizado exitosamente en:\n  {master_docx}")
    except PermissionError:
        print(f"[AVISO] '{master_docx}' está abierto en Microsoft Word.")
        backup_master = os.path.join(compilado_dir, "El_Camino_del_Reencarnado_temp.docx")
        create_novel_document(titulo, subtitulo, autor, capitulos, backup_master)
        print(f"[OK] Guardado en copia auxiliar:\n  {backup_master}")


def ejecutar():
    print("===============================================================")
    print(f" AGREGAR / ACTUALIZAR: {CAPITULO['header']} - {CAPITULO['title']}")
    print("===============================================================")
    for base in [CURRENT_DIR, NOVELA_DIR, PROYECTO_DIR]:
        if os.path.exists(base):
            print(f"[*] Procesando en: {base}")
            guardar_capitulo_individual(CAPITULO, base)
            compilar_libro_completo(base)
    print("===============================================================")
    print(" ¡PROCESO COMPLETADO EXITOSAMENTE!")
    print("===============================================================")

if __name__ == "__main__":
    ejecutar()
