# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_10 = {
    "header": "Capítulo 10",
    "title": "La Unión de la Clase F",
    "images": [
        (os.path.join(NOVELA_DIR, "imagenes", "subaru_natsuki.png"), "Subaru Natsuki — Logística y Apoyo"),
        (os.path.join(NOVELA_DIR, "imagenes", "erza_scarlet.png"), "Erza Scarlet — Disciplina y Vanguardia de Acero"),
        (os.path.join(NOVELA_DIR, "imagenes", "yami_sukehiro.png"), "Yami Sukehiro — Fuerza y Combate Pragmático"),
        (os.path.join(NOVELA_DIR, "imagenes", "akame.png"), "Akame — Sigilo, Rastreo y Precisión"),
        (os.path.join(NOVELA_DIR, "imagenes", "rukia_renji.jpg"), "Rukia Kuchiki y Renji Abarai — Asalto y Soporte Arcano")
    ],
    "paragraphs": [
        "El domingo previo a la partida, el patio de armas secundario se convirtió en una colmena de actividad.",
        "Para sobrevivir una semana entera en el bosque y explorar un calabozo subterráneo, no bastaba con que nuestro pequeño grupo estuviera coordinado; necesitábamos que toda la Clase F-2 funcionara como un solo cuerpo.",
        "Sin embargo, varios de nuestros compañeros seguían manteniéndose apartados en los rincones del patio por distintos motivos: algunos por desconfianza, otros por timidez o por un orgullo solitario forjado en pasados difíciles.",
        "Fue en ese momento cuando Ken y Kaoruko demostraron por qué eran el corazón de nuestra clase.",
        "* * *",
        "—¡No podemos irnos a la boca del lobo si ni siquiera sabemos los nombres de quienes nos van a cuidar la espalda! —bramó Ken con su habitual entusiasmo, remangándose la chaqueta y dando saltos en el sitio.",
        "—Tiene razón, Ken-kun —lo secundó Kaoruko, sosteniendo una gran canasta llena de panecillos calientes, carne seca sazonada y té aromático—. Nadie se negará si nos acercamos con una sonrisa... ¡y con algo de comida!",
        "Ken fue directo hacia el muro del fondo, donde un chico de chándal negro y naranja con el pelo alborotado revisaba torpemente una cantimplora.",
        "—¡Oye, tú! ¡Subaru Natsuki, verdad! ¡Deja esa cantimplora y ven a la mesa grande! —le gritó Ken pasándole un brazo pesado por los hombros.",
        "—¡A-Ah! ¡Oye, despacio, salvaje, casi me dislocas la clavícula! —se quejó Subaru Natsuki con una mueca cómica, aunque al ver la enorme sonrisa de Ken y a Kaoruko ofreciéndole un panecillo, se rascó la nuca y suspiró aliviado—. Bueno... supongo que la logística del campamento no se va a organizar sola. Si necesitan a alguien con sentido común, cuenten conmigo.",
        "Mientras tanto, Kaoruko se acercó con paso dulce hacia una chica pelirroja de mirada severa que pulía su armadura plateada con un paño de cuero.",
        "—Disculpe... señorita Erza —dijo Kaoruko con timidez, ofreciéndole una taza humeante—. Preparamos té para la reunión... nos encantaría contar con su experiencia para organizar las guardias.",
        "Erza Scarlet levantó la vista. Su mirada afilada se suavizó al percibir la genuina amabilidad de Kaoruko. Tomó la taza con ambas manos, haciendo una inclinación formal de cabeza.",
        "—Acepto con gusto —respondió Erza con voz firme y marcial—. En una expedición, la disciplina y las formaciones lo son todo. Me aseguraré de que nadie rompa filas.",
        "A pocos metros, apoyado contra un pilar de piedra con los brazos cruzados y una katana al cinto, Yami Sukehiro observaba la escena con aire relajado y una brizna de hierba entre los dientes.",
        "—Vaya alboroto tienen armado —dijo Yami con su voz grave y rasposa cuando me acerqué a él—. Aunque supongo que un poco de ruido es mejor que morir desangrado en una cueva.",
        "—Ven a la mesa, Yami —le dije con calma—. Tus instintos en combate nos vendrán bien.",
        "—Está bien, grandullón —asintió Yami con una sonrisa ladeada, ajustándose la vaina de su espada—. Pero no esperen discursos bonitos de mi parte. Cuando una bestia salte del lodo, solo hay una regla: superar los límites y partirla en dos.",
        "La comida de Kaoruko también atrajo de inmediato a Akame. La chica de larga melena negra y ojos rojos, que hasta ese momento afilaba su hoja en silencio en la sombra, apareció junto a la mesa en un abrir y cerrar de ojos, con la mirada fija en los platos de carne.",
        "—Huele delicioso —dijo Akame con su tono calmado y directo, masticando un trozo con evidente satisfacción—. Puedo encargarme del rastreo silencioso de bestias y del despiece de comida en el bosque.",
        "Por último, se sumaron Rukia Kuchiki y Renji Abarai. Renji llegó quejándose ruidosamente con su cabello rojo en puntas y sus tatuajes al aire, mientras Rukia le daba un pisotón por sus malos modales antes de presentarse formalmente con su porte sereno de espadachina.",
        "* * *",
        "Con todos sentados alrededor de los bancos corridos, la mesa de la Clase F-2 cobró una energía arrolladora.",
        "Éramos un grupo dispar como ninguno: Seier Silva, Kaoruko Waguri, Subaru Hoshina, Ken, los seis Hijos de Calydon liderados por Caesar y Lighter, junto a Subaru Natsuki, Erza Scarlet, Yami Sukehiro, Akame, Rukia y Renji.",
        "Desplegamos el mapa del Bosque de las Sombras Bajas sobre los tablones.",
        "—Muy bien, escuchen todos —dijo Caesar King golpeando la madera con los nudillos para pedir orden—. La vanguardia de contención estará a cargo de Seier, Lighter, Renji y yo. Nosotros absorbemos el choque frontal.",
        "—Erza y Yami pueden cubrir los flancos de asalto pesado con sus filos —propuse, señalando las entradas angostas del valle—. Akame y Pulchra se encargarán del reconocimiento avanzado entre los árboles.",
        "—Subaru Natsuki y yo nos aseguraremos del inventario, las raciones y las rutas de escape —intervino Subaru Hoshina con su habitual libreta de notas.",
        "—Y en la retaguardia estaremos nosotros —dijo Kaoruko con una sonrisa decidida—. Con mi magia de agua, las barreras de Rukia y la ayuda médica, nadie se quedará atrás.",
        "—¡Ese es un plan digno de un gremio legendario! —bramó Ken alzando el puño con entusiasmo.",
        "Las risas, las bromas y las discusiones tácticas se extendieron hasta el atardecer. Las dudas y la distancia que antes separaban a la clase se habían evaporado por completo.",
        "* * *",
        "El lunes por la mañana, los portones occidentales de la academia se abrieron bajo una suave brisa fresca.",
        "Toda la Clase F-2 avanzó en formación, con las mochilas ajustadas a la espalda, las armas aseguradas al cinto y la determinación marcada en el rostro. Los nobles de las clases superiores nos miraban desde los balcones con desdén, convencidos de que fracasaríamos en el barro.",
        "No tenían idea de lo equivocados que estaban.",
        "Miré a mis compañeros a la izquierda y a la derecha. Éramos los rechazados, los olvidados y los provincianos, pero caminábamos juntos como una verdadera manada.",
        "Apreté el medallón de la Casa Silva bajo mi camisa, respiré el olor a pino del sendero y di el primer paso hacia la penumbra del bosque.",
        "La primera gran prueba de la Clase F había comenzado."
    ]
}

def compilar_capitulo_10():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10]
    
    # 1. Guardar texto plano del Capítulo 10
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "10_Capitulo_10.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_10['header'].upper()}: {CAPITULO_10['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_10['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 10 con todas las ilustraciones
    docx_cap10 = os.path.join(NOVELA_DIR, "capitulos", "10_Capitulo_10.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_10['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_10['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    # Insertar imágenes de los compañeros integrados
    for img_path, caption in CAPITULO_10.get("images", []):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            r_img = p_img.add_run()
            r_img.add_picture(img_path, width=Inches(3.2))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(14)
            r_cap = p_cap.add_run(f"Ilustración: {caption}")
            r_cap.font.name = "Georgia"
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(120, 100, 140)

    for p_text in CAPITULO_10['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap10)
    print(f"[OK] Word individual Capítulo 10 guardado en: {docx_cap10}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 10)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_10()
