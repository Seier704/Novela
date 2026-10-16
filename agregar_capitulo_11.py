# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_11 = {
    "header": "Capítulo 11",
    "title": "La Forja en las Sombras Bajas",
    "paragraphs": [
        "El Bosque de las Sombras Bajas no perdonaba a los desprevenidos.",
        "Apenas cruzamos el linde de los árboles, la luz del mediodía quedó filtrada por una cúpula asfixiante de copas milenarias. El aire olía a tierra húmeda, musgo y el sutil hedor ferroso que dejaban las bestias corrompidas por el miasma.",
        "Los primeros dos días fueron una prueba silenciosa de resistencia y logística.",
        "Subaru Natsuki demostró de inmediato por qué su presencia era vital: mientras los demás pensábamos solo en golpear, él se convirtió en el cerebro organizativo del campamento. Junto a Subaru Hoshina, trazó un perímetro defensivo junto a una curva del río, asignó turnos de guardia rotativos y calculó al gramo las provisiones.",
        "—¡Si dejamos la comida en el suelo, las hormigas carroñeras se la comerán en media hora! —regañaba Natsuki, colgando sacos en las ramas altas con cuerdas—. ¡Y tú, Ken, deja de usar las antorchas para encender ramas verdes o nos vas a ahumar a todos!",
        "—¡Jajaja, pero si hace frío, hombre de chándal! —respondía Ken riendo mientras armaba las estacas de las tiendas con golpes limpios de sus manos.",
        "La comida tampoco fue un problema gracias a Akame y Pulchra. Ambas se internaban en la espesura como dos sombras mudas y regresaban al campamento cargando con jabalíes de roca y ciervos de cuernos de bronce. Kaoruko se encargaba de encender una olla comunitaria gigante, sazonando los guisos con hierbas silvestres que hacían que hasta el severo Yami Sukehiro se sentara en silencio a repetir tres platos.",
        "* * *",
        "La noche del martes fue la primera en que el bosque nos mostró los dientes.",
        "Me encontraba montando la primera guardia junto a Kaoruko en un saliente rocoso que dominaba el claro del campamento. El fuego central ardía en rescoldos bajos para no delatar nuestra posición.",
        "Kaoruko estaba envuelta en una manta gruesa, con las manos alrededor de una taza de té tibio. Sus enormes ojos azules miraban la oscuridad con una mezcla de respeto y nerviosismo.",
        "—Seier-kun... —susurró ella, rompiendo el susurro del viento entre las hojas—. ¿Alguna vez sientes miedo?",
        "—Todo el mundo siente miedo —respondí con mi tono tranquilo, sentado a su lado con la espalda apoyada en una peña—. En mi tierra natal... y en la frontera sur con mis padres, aprendí que el miedo no es malo. Es lo que te mantiene alerta para no cometer estupideces. Lo que importa es lo que decides hacer con el cuerpo cuando el miedo te aprieta el pecho.",
        "Kaoruko giró su rostro hacia mí. La tenue luz de las brasas dibujaba destellos cálidos sobre sus mejillas suaves.",
        "—Cuando estoy cerca de ti... casi no lo siento —admitió en voz muy baja, bajando un instante la vista con un rubor que apenas pudo disimular—. Tienes una forma de estar de pie que hace que todo parezca que saldrá bien.",
        "—Es por los noventa kilos —bromeé con una leve sonrisa—. Cuesta moverme del sitio.",
        "Ella soltó una risita dulce y cantarina, apoyando suavemente su hombro contra mi brazo por un segundo antes de que un aullido gutural y distante cortara la noche.",
        "* * *",
        "El miércoles por la madrugada se produjo el primer gran choque.",
        "Una jauría de más de treinta Lobos de Miasma —bestias con pelaje erizado de púas negras y ojos amarillentos como ascuas podridas— intentó emboscar el campamento aprovechando la niebla del río.",
        "No hubo pánico.",
        "—¡Alarma! ¡Formación de cuña! —la voz de trueno de Erza Scarlet resonó al instante. Su espada desenvainada destelló con un resplandor plateado mientras cortaba en dos al primer lobo que saltó sobre las estacas.",
        "Caesar King, Lighter, Renji y yo saltamos al frente de la empalizada.",
        "—¡Sostengan la línea, idiotas! —rugió Renji, soltando estocadas feroces con su hoja de colmillos dentados.",
        "Un lobo gigante, el doble de grande que el resto, se abalanzó directo hacia donde estaba Kaoruko preparando los círculos de curación. Di dos zancadas pesadas y me interpuse en su trayectoria. Planté los pies con la firmeza del elemento tierra: el choque del animal contra mi pecho fue como el impacto de un carromato, pero mis piernas no cedieron ni un milímetro.",
        "Enganché su cuello con un clinch de hierro, giré el torso y descargué un rodillazo ascendente directamente en su esternón.",
        "¡CRACK!",
        "El aire quemó mis nudillos. Por primera vez en combate real, el fuego brotó con fluidez desde mi estómago hacia mis manos: rematé con un codazo envuelto en llamas en la sien de la bestia, calcinando el miasma en una detonación sorda.",
        "—¡Buen golpe, Silva! —bramó Caesar, derribando a otros dos con un barrido de su escudo militar.",
        "Por los flancos, Yami se movía con una calma aterradora: cada tajo de su katana dejaba un rastro de vacío oscuro que desmembraba a los engendros antes de que pudieran gruñir. Akame remataba a los que intentaban huir con puñaladas quirúrgicas al corazón.",
        "En menos de diez minutos, la jauría entera fue aniquilada. Ningún estudiante de la F-2 sufrió heridas de gravedad gracias a las barreras de hielo de Rukia y la rápida intervención sanadora de Kaoruko y Subaru Hoshina.",
        "* * *",
        "El jueves al amanecer encontramos por fin el objetivo principal: la entrada al Calabozo de Piedra Gris.",
        "Era una grieta monumental abierta en las raíces de una colina, flanqueada por pilares de piedra en ruinas tallados con runas arcanas. Un aire gélido y cargado de maná emanaba de las profundidades.",
        "—Las lecturas indican trampas de contrapeso en los primeros tres niveles —informó Lucy, leyendo un pergamino con el ceño fruncido.",
        "—Dejen que Pulchra y yo vayamos en punta —propuso Akame con voz serena—. Desactivaremos los mecanismos antes de que el grupo pesado avance.",
        "El avance por los pasillos subterráneos fue una obra maestra de coordinación. Natsuki marcaba las paredes con tiza luminosa para no perder la ruta; Burnice iluminaba los techos con llamaradas controladas que impedían que los murciélagos de sombra cayeran sobre la retaguardia; y Piper, pese a su cara de sueño, usaba su enorme llave mecánica para trabar compuertas de piedra que amenazaban con cerrarse a nuestro paso.",
        "Hasta que llegamos a la sala del jefe en el nivel más profundo.",
        "* * *",
        "La caverna central era gigantesca, sostenida por estalagmitas que brillaban con un musgo fosforescente.",
        "En el centro dormía el guardián de la mazmorra: un Gólem Acorazado de Miasma, una mole de roca negra y metal oxidado de cuatro metros de altura, con un núcleo escarlata latiendo en mitad de su pecho como un corazón podrido.",
        "Al sentir nuestras presencias, el coloso despertó con un bramido que hizo temblar el techo de la cueva.",
        "—¡Es inmune al corte normal! —advirtió Erza tras probar su hoja contra el hombro de piedra, sacando apenas chispas—. ¡Necesitamos fisurar el caparazón antes de atacar el núcleo!",
        "—¡Déjenmelo a mí y al cabeza azul! —gritó Burnice, lanzando un chorro de combustible que Ken encendió con sus puños en llamas. El fuego envolvió las piernas de roca del gólem, calentando la piedra al rojo vivo.",
        "—¡Ahora, Rukia, enfríalo de golpe! —ordenó Natsuki desde la línea de mando.",
        "Rukia alzó su espada blanca: un torrente de escarcha helada cubrió la roca incandescente. La física hizo el resto: el choque térmico resquebrajó la coraza del monstruo con crujidos secos.",
        "—¡Vanguardia, a la carga! —rugió Caesar King.",
        "Lighter esquivó un manotazo descomunal del gólem con un balanceo perfecto de cintura y le destrozó la rodilla con un gancho demoledor al tendón de roca. Renji y Caesar le abrieron una brecha en el costado.",
        "El gólem tambaleó, alzando el puño izquierdo para aplastar a Kaoruko en la retaguardia.",
        "No lo dudé un segundo.",
        "Exploté hacia adelante con el paso más rápido de Jeet Kune Do que jamás había ejecutado. Salté sobre una piedra saliente y me lancé directamente hacia el pecho agrietado de la mole.",
        "Concentré toda la energía: el peso telúrico de la tierra en mi hombro y la detonación del fuego en mi puño derecho cerrado.",
        "—¡ARDE Y ROMPE! —rugí.",
        "¡BOOOOOOM!",
        "Mi puñetazo impactó en la fisura del pecho con la fuerza de un meteorito. El fuego explotó hacia el interior de la roca al mismo tiempo que la onda de choque de tierra pulverizaba la estructura interna.",
        "El núcleo escarlata estalló en mil pedazos de cristal inerte.",
        "El gólem se detuvo en seco, crujió de arriba abajo y se desplomó contra el suelo en una montaña de cascotes silenciosos.",
        "Silencio.",
        "Y luego, un grito unísono de júbilo y victoria que retumbó por todo el calabozo.",
        "* * *",
        "El domingo al atardecer, la expedición llegó a su fin.",
        "Los portones principales de la Gran Academia de Solaria estaban atestados de profesores, examinadores y alumnos de las clases altas esperando ver qué quedaba de los marginados de la F-2.",
        "Las puertas de hierro chirriaron al abrirse.",
        "Nadie volvió rengueando. Nadie volvió quejándose. Toda la Clase F-2 cruzó la entrada en formación militar perfecta, con mochilas cargadas de núcleos mágicos, pieles de lobo y cristales del calabozo.",
        "Al frente íbamos Caesar King y yo, cubiertos de polvo y hollín, pero con las cabezas bien altas. Detrás de nosotros marchaban Lighter, Erza, Yami con su katana al hombro, Akame, Ken riendo a carcajadas con Renji, y en el centro, protegidas y sonrientes, Kaoruko y Subaru flanqueadas por Natsuki y los Hijos de Calydon.",
        "Los nobles de las terrazas superiores guardaron un silencio sepulcral.",
        "Alcé la vista hacia el balcón de la Clase A.",
        "Allí estaba Elysea de Vane. Apoyada en la barandilla de mármol blanco, con su porte aristocrático, me miraba fijamente. Esta vez no había burla ni desdén en sus ojos azul zafiro; había asombro puro, respeto indisimulado y un brillo que quemaba más que el hielo.",
        "Le sostuve la mirada durante dos segundos, sonreí con calma y seguí marchando junto a mis camaradas.",
        "La Clase F ya no era el fondo del barril.",
        "Ahora éramos la fuerza más temible de toda la academia."
    ]
}

def compilar_capitulo_11():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11]
    
    # 1. Guardar texto plano del Capítulo 11
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "11_Capitulo_11.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_11['header'].upper()}: {CAPITULO_11['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_11['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 11
    docx_cap11 = os.path.join(NOVELA_DIR, "capitulos", "11_Capitulo_11.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_11['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_11['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_11['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap11)
    print(f"[OK] Word individual Capítulo 11 guardado en: {docx_cap11}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 11)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_11()
