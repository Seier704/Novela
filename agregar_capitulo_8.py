# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_8 = {
    "header": "Capítulo 8",
    "title": "Los Hijos de Calydon y el Golpe del Campeón",
    "illustration": os.path.join(NOVELA_DIR, "imagenes", "hijos_de_calydon.png"),
    "illustration_caption": "Ilustración: Los Hijos de Calydon (Clase F-2)",
    "paragraphs": [
        "La segunda semana en la academia comenzó con un cambio de ritmo en las lecciones.",
        "El profesor Vaneck empezó a mezclar las clases de teoría mágica con simulaciones de combate en equipo y tácticas de supervivencia en grupo. En un aula como la F-2, donde casi todos éramos marginados o plebeyos con distintas habilidades, aprender a coordinarse era una cuestión de supervivencia básica.",
        "Fue durante esas prácticas que me acerqué a conocer a fondo al grupo más peculiar y ruidoso de toda la clase: una banda de seis estudiantes que siempre se movían en bloque y actuaban como una verdadera familia.",
        "Se hacían llamar los Hijos de Calydon.",
        "* * *",
        "Aprovechamos un descanso entre las sesiones de estrategia para sentarnos en las gradas de madera del coliseo secundario.",
        "Al frente del grupo estaba Caesar King, una chica imponente de mirada fiera, cabello corto verde ceniza y una pequeña corona oscura entre los mechones que le daba un aire de reina rebelde. A su lado estaban Luciana de Montefio, a quien llamaban Lucy, una rubia con gorra militar y coleta que no paraba de dar órdenes con voz chillona; Piper Wheel, una chica bajita y somnolienta que bostezaba apoyada en una llave mecánica gigante; Burnice, una rubia hiperactiva con gafas rojas que siempre olía a combustible y sonreía con ganas de hacer explotar algo; y Pulchra, una ágil semi-humana con orejas y cola de zorro que vigilaba todo en silencio tras una máscara de tela.",
        "Y finalmente estaba el sexto miembro: Lighter.",
        "Lighter era un muchacho alto y atlético, de cabello oscuro algo despeinado, gafas de sol oscuras que rara vez se quitaba, una bufanda roja al cuello y una chaqueta de cuero negra con hombreras metálicas. Caminaba con las manos en los bolsillos y una soltura que gritaba experiencia.",
        "—Así que tú eres Seier Silva —dijo Caesar King con una sonrisa amplia y desafiante, cruzándose de brazos frente a mí—. He visto cómo te mueves. Tienes madera de líder de vanguardia, grandullón.",
        "—Caesar es la que manda aquí —aclaró Lucy con las manos en las caderas, mirándome con orgullo—. Venimos todos de la misma aldea en los páramos de Calydon. Éramos un grupo de pequeños aventureros allá afuera antes de que nos enviaran a esta academia.",
        "—Una aldea dura, supongo —dije, asintiendo con respeto.",
        "—Lo bastante dura como para que nadie sobreviva siendo un blando —intervino Lighter, dando un paso al frente con una media sonrisa tranquila—. Y hablando de dureza... veo que tú tampoco usas varitas ni espadas ceremoniales, Silva. Te gusta resolver las cosas con los puños.",
        "—Es lo más directo y lo que mejor funciona —respondí mirándole las manos vendadas—. ¿Tú también peleas desarmado?",
        "—Boxeo puro —dijo Lighter, tronándose los nudillos con un chasquido seco—. Nada de patadas acrobáticas ni poses raras. Solo velocidad, esquiva y manos pesadas. ¿Qué dices? ¿Hacemos un pequeño sparring amistoso antes de que vuelva el instructor?",
        "—¡Oh, esto se va a poner bueno! —gritó Burnice dando saltitos de emoción—. ¡Cien monedas a Lighter!",
        "—No apuestes con dinero que no tienes, Burnice... —bostezó Piper desde el banco.",
        "—Tengan cuidado los dos... —murmuró Kaoruko desde atrás, juntando las manos con su habitual preocupación.",
        "* * *",
        "Nos colocamos en el centro de la arena circular de tierra.",
        "Lighter se acomodó las gafas de sol, adelantó el pie izquierdo y levantó la guardia clásica de boxeador: puños pegados a las mejillas, codos cerrados protegiendo las costillas y la barbilla hundida en el pecho. Su postura se balanceaba con un ritmo suave e hipnótico.",
        "Yo bajé el centro de gravedad, abrí ligeramente los pies en guardia de Muay Thai y fijé la mirada en su pecho.",
        "—Cuando quieras, Silva —dijo Lighter.",
        "Exploté hacia adelante buscando acortar la distancia con un paso de intercepción de Jeet Kune Do. Lancé un directo de derecha al mentón con todo el peso de mis noventa kilos.",
        "El golpe cortó el aire con fuerza, pero Lighter no estaba allí.",
        "Con un movimiento de cintura fluido y milimétrico, dobló el torso hacia un lado, dejando que mi puño pasara a milímetros de su oreja. En el mismo segundo de la esquiva, su mano izquierda salió disparada como un pistón:",
        "¡PAF!",
        "Un jab seco y cegador impactó directo en mi pómulo, seguido de inmediato por un gancho al cuerpo que golpeó mis costillas flotantes antes de que pudiera cerrar el codo.",
        "El dolor fue agudo, pero apreté los dientes y no retrocedí. Tiré una patada baja con la tibia de hierro hacia su pierna adelantada, buscando quebrar su juego de pies.",
        "Lighter leyó la trayectoria al instante: levantó la rodilla para amortiguar el impacto, giró sobre sus talones y soltó una combinación relámpago de tres golpes: mentón, sien y plexo.",
        "Logré meter los antebrazos reforzados con maná para bloquear los dos primeros, pero la velocidad de sus manos era sobrehumana. Era como intentar atajar una ráfaga de piedras disparadas por una catapulta.",
        "Vi un hueco y entré con el hombro por delante, conectándole un golpe corto con la izquierda en la mandíbula que le hizo volar las gafas oscuras y lo obligó a retroceder dos pasos, sacudiendo la cabeza con sorpresa.",
        "—¡Buena mano! —reconoció Lighter con los ojos brillantes, frotándose la barba incipiente—. Tienes el impacto de una maza.",
        "—Y tú te mueves como una avispa —respondí, respirando hondo.",
        "Lighter aceleró el ritmo. Ya no se contuvo: combinó esquivas por debajo de mis ganchos con fintas de hombro que me hicieron perder la trayectoria. En una fracción de segundo, esquivó mi directo con un cabeceo perfecto, entró en mi punto ciego y descargó un golpe ascendente al estómago seguido de un crochet demoledor a la mandíbula.",
        "El impacto me sacudió el cerebro. Mis piernas se desconectaron por un instante y caí sobre una rodilla en la arena, apoyando la mano en el suelo mientras intentaba recuperar el aire.",
        "—¡Tiempo! —anunció Caesar King con una palmada sonora.",
        "* * *",
        "Lighter respiraba agitado, con un moretón visible en la mandíbula y el labio partido por el golpe que le alcancé a conectar. Se agachó, recogió sus gafas de sol del suelo, se las puso de nuevo y me extendió su mano enguantada con una sonrisa de oreja a oreja.",
        "—Maldita sea, Silva... casi me arrancas la cabeza con ese zurdazo —dijo Lighter, ayudándome a ponerme en pie con un tirón firme—. Si tuvieras un poco más de velocidad de pies, la historia habría sido otra. Tu pegada es aterradora.",
        "—Tu velocidad y tus esquivas están en otro nivel —admití, sacudiéndome el polvo y frotándome la barbilla—. Se nota que no aprendiste eso en un aula.",
        "—Años peleando contra bestias y bandidos en los caminos de Calydon —dijo Caesar acercándose y dándome una palmada amistosa en la espalda que casi me mueve el esqueleto—. Lighter es nuestro mejor peleador por una razón. Pero aguantarle el ritmo y conectarle un golpe limpio en su primer combate... te has ganado mi respeto, Seier.",
        "—¡Yo sabía que Lighter ganaría! ¡Pero estuviste increíble, chico nuevo! —exclamó Burnice aplaudiendo con entusiasmo.",
        "Kaoruko se acercó corriendo con una toallita húmeda y su cantimplora, examinándome la cara con los ojos muy abiertos.",
        "—¡Seier-kun, tienes el pómulo morado! Déjame curarte rápido...",
        "—Estoy bien, Kaoruko. Solo fue un buen intercambio —le dije con calma, dejándome aplicar su magia fresca sobre la piel.",
        "Lighter y Ken se acercaron también, empezando a comparar técnicas y hablar sobre ángulos de golpeo con la emoción propia de quienes viven para el combate.",
        "En una academia llena de nobles que nos miraban con desprecio, la Clase F-2 estaba reuniendo a los mejores guerreros que el reino jamás vio venir. Y con los Hijos de Calydon de nuestro lado, el tablero de juego acababa de ponerse mucho más interesante."
    ]
}

def compilar_capitulo_8():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8]
    
    # 1. Guardar texto plano del Capítulo 8
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "08_Capitulo_8.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_8['header'].upper()}: {CAPITULO_8['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_8['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 8 con ilustración
    docx_cap8 = os.path.join(NOVELA_DIR, "capitulos", "08_Capitulo_8.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_8['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_8['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)
    
    if os.path.exists(CAPITULO_8['illustration']):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(15)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        # Ancho adecuado para imagen panorámica de 6 personajes
        r_img.add_picture(CAPITULO_8['illustration'], width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(20)
        r_cap = p_cap.add_run(CAPITULO_8['illustration_caption'])
        r_cap.font.name = "Georgia"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(120, 100, 140)

    for p_text in CAPITULO_8['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap8)
    print(f"[OK] Word individual Capítulo 8 guardado en: {docx_cap8}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 8)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_8()
