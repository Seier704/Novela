# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_16 = {
    "header": "Capítulo 16",
    "title": "La Batalla de los Libros",
    "paragraphs": [
        "Las dos semanas previas a los exámenes convirtieron la sala común de la Clase F-2 en un campo de batalla infinitamente más sangriento que el Bosque de las Sombras Bajas.",
        "Los pupitres se alinearon formando una mesa comunal atestada de atlas amarillentos, pergaminos de runas y tazas vacías de café amargo. Al frente del régimen de estudio, las chicas tomaron el mando con mano de hierro.",
        "—¡Ken! ¡Si vuelves a quedarte dormido sobre el mapa del Río de Maná, juro por mi espada que te clavo las mangas a la mesa! —bramó Erza Scarlet, golpeando una regla de madera gruesa contra el tablón con el porte de un general de asedio.",
        "—¡P-Pero Erza, los nombres de los reyes fundadores suenan todos iguales! —gimoteaba Ken con los ojos enrojecidos y el flequillo azul pegado a la frente de tanto sudar.",
        "—¡Cállate y memoriza las dinastías del sur! —le espetó Lucy desde la otra punta, mientras obligaba a Burnice a repetir proporciones químicas de barreras mágicas antes de que le permitiera tocar sus propulsores de fuego.",
        "A pocos metros, Renji intentaba escapar sigilosamente por una ventana entreabierta, pero un certero manotazo de Rukia con un tomo de derecho feudal lo regresó de espaldas a su asiento entre maldiciones.",
        "Incluso Yami estaba sentado a regañadientes con los brazos cruzados, mascando su cigarro apagado mientras Natsuki —quien resultó tener una memoria prodigiosa para las fechas y la geografía— le explicaba las rutas comerciales del reino dibujando garabatos cómicos en una pizarra portátil.",
        "—Mira, Yami, imagínate que este reino vecino es un tipo pesado al que quieres cortar en dos... sus defensas mágicas están concentradas aquí en el paso de montaña —explicaba Natsuki con entusiasmo pedagógico.",
        "—Hmm... visto así tiene más sentido —admitió Yami rascándose la nuca—. Buen trabajo, novato.",
        "* * *",
        "En el rincón más apartado y tranquilo de la sala, Kaoruko y yo compartíamos una pequeña mesa iluminada por un candil de aceite.",
        "A diferencia de los gritos y amenazas del resto del salón, estudiar con ella era un remanso de paz. Se sentaba a mi lado, tan cerca que nuestras rodillas se rozaban bajo el pupitre, señalando las líneas de los pergaminos con la punta de su pluma de tinta.",
        "—El flujo de tierra no se proyecta hacia afuera como el fuego, Seier-kun —me explicaba con su voz dulce y paciente, inclinándose sobre el texto—. Tienes que memorizar la runa de contención telúrica... mira cómo se trazan estas tres espirales.",
        "Sentía el calor suave de su piel y el aroma fresco a lirios que desprendía su cabello negro cada vez que se movía. Cuando levantaba la vista para responder a sus preguntas sobre táctica militar o historia fronteriza, nuestros ojos se encontraban a escasos centímetros.",
        "—La runa de contención requiere tres puntos de anclaje —respondí mirándola fijamente—, igual que cuando plantas los pies antes de recibir una embestida.",
        "Kaoruko parpadeó sorprendida, y una sonrisa radiante iluminó sus mejillas sonrojadas.",
        "—¡Exacto! ¡Lo entendiste a la primera! —exclamó juntando las manos, sacando un pequeño pastelillo de miel de su canasta y poniéndomelo en la palma de la mano con picardía—. Premio por la respuesta correcta.",
        "—A este paso voy a terminar aprobando con honores solo por la comida, Kaoruko —bromeé dándole un bocado.",
        "—¡Esa es la idea! Quiero que vayamos todos juntos a esas aguas termales —dijo ella con una timidez sincera, bajando la vista hacia sus apuntes con una sonrisa que me quitó de golpe todo el cansancio.",
        "Gracias a mi memoria de mi vida anterior y a las explicaciones minuciosas de Kaoruko, los conceptos abstractos de maná y las leyes feudales empezaron a encajar en mi cabeza como las piezas de una armadura bien forjada.",
        "* * *",
        "El día del examen semestral, el Gran Salón de la Academia parecía un tribunal de juicio final.",
        "Filas interminables de pupitres individuales se extendían bajo los arcos góticos de la nave central, vigilados por diez archimaestres con túnicas ceremoniales y varitas de detección de trampas arcanas flotando en el aire. Cientos de aspirantes de todas las clases llenaban los asientos en un silencio tenso.",
        "Al entrar la Clase F-2, los murmullos de las Clases A y B no tardaron en escucharse entre risitas burlonas:",
        "—Miren a los brutos de la F... deben estar rezando para no entregar la hoja en blanco.",
        "—Agradezcan si logran escribir su propio nombre.",
        "Ignoramos las provocaciones.",
        "Al sentarme en mi puesto asignado, alcancé a ver a Elysea de Vane tres filas más adelante. La elfa de hielo me miró de reojo con sus penetrantes ojos azul zafiro, arqueando una ceja plateada con una mezcla de curiosidad y expectativa desafiante. Le devolví una mirada tranquila y serena antes de centrarme en mi escritorio.",
        "—¡El tiempo comienza ahora! —anunció el juez principal con un golpe de campana.",
        "Cien plumas comenzaron a rasguear el papel al unísono.",
        "Las tres horas siguientes fueron una prueba de fuego mental.",
        "Las preguntas eran complejas: fórmulas de conversión de maná elemental, historia de las guerras contra el Abismo y análisis táctico de formaciones de asedio. Vi a Ken sudar frío mientras se mordía el nudillo de la mano izquierda, a Renji rascándose la cabeza con la pluma torcida y a Natsuki escribiendo párrafos enteros a toda velocidad.",
        "Por mi parte, encaré el examen con la misma frialdad con la que cerraba la distancia en un combate: recordé cada explicación paciente de Kaoruko, vinculé las tácticas militares con lo vivido en el bosque y redacté respuestas directas, pragmáticas y sin adornos innecesarios.",
        "—¡Tiempo! ¡Plumas en la mesa! —decretó el juez.",
        "Un suspiro colectivo de alivio y agotamiento retumbó en la bóveda.",
        "Ken se dejó caer bocarriba en el asiento como si le hubieran dado un martillazo; Renji soltó un bufido de aire caliente; pero nadie, absolutamente nadie de la F-2, había dejado su hoja vacía.",
        "Al salir al patio bañado por la luz del atardecer, Kaoruko se acercó a mi lado con una sonrisa esperanzada.",
        "—¿Cómo te fue, Seier-kun?",
        "—Contesté todo —le respondí, chocando su puño suavemente con los nudillos—. Gracias a ti.",
        "Los dados estaban tirados. Y la Clase F acababa de demostrar que no solo sabía pelear con los puños, sino también con el cerebro."
    ]
}

def compilar_capitulo_16():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    from agregar_capitulo_12 import CAPITULO_12
    from agregar_capitulo_13 import CAPITULO_13
    from agregar_capitulo_14 import CAPITULO_14
    from agregar_capitulo_15 import CAPITULO_15
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12, CAPITULO_13, CAPITULO_14, CAPITULO_15, CAPITULO_16]
    
    # 1. Guardar texto plano del Capítulo 16
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "16_Capitulo_16.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_16['header'].upper()}: {CAPITULO_16['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_16['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 16
    docx_cap16 = os.path.join(NOVELA_DIR, "capitulos", "16_Capitulo_16.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_16['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_16['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_16['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap16)
    print(f"[OK] Word individual Capítulo 16 guardado en: {docx_cap16}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 16)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_16()
