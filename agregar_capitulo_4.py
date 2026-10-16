# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_4 = {
    "header": "Capítulo 4",
    "title": "La Chica de Mirada Fría",
    "illustration": os.path.join(NOVELA_DIR, "imagenes", "subaru_hoshina.png"),
    "illustration_caption": "Ilustración: Subaru Hoshina (Clase F-2)",
    "paragraphs": [
        "Los primeros tres días en la academia pasaron con bastante rapidez y no pocas novedades.",
        "La rutina de la Clase F-2 se fue consolidando en un ritmo exigente: por las mañanas teníamos sesiones de acondicionamiento físico militar y teoría estratégica; por las tardes, adaptación al flujo de maná y lecciones de historia.",
        "Durante esos entrenamientos, no tardé en notar las fortalezas y debilidades de mis compañeros, y en especial las de Kaoruko.",
        "En lo físico, a la pobre chica le costaba horrores. Las carreras de resistencia la dejaban sin aliento a mitad del circuito y levantar peso le provocaba temblores en los brazos. Sin embargo, cuando llegaba la hora de canalizar magia, su desempeño era notable. Tenía una afinidad limpia con el elemento agua, pero en lugar de proyectar lanzas o látigos acuáticos, se volcaba por completo en los hechizos de sanación.",
        "* * *",
        "Esa tarde, tras terminar una agotadora práctica de campo, nos sentamos bajo la sombra de un gran sauce en el patio oeste.",
        "Kaoruko estaba sentada en la hierba secándose el sudor de la frente con una toalla pequeña, respirando aún con agitación mientras sostenía su cantimplora.",
        "—Buen esfuerzo hoy —le dije, sentándome a su lado con las piernas cruzadas.",
        "—G-Gracias, Seier-kun... aunque siento que mis piernas van a derretirse —respondió ella con una risita cansada, apoyando la espalda contra el tronco del árbol—. De verdad soy un desastre corriendo.",
        "—Solo te falta costumbre. En cambio, tus conjuros de agua salieron bastante bien. La esfera de restauración que hiciste duró el doble que la de los demás.",
        "Kaoruko sonrió, satisfecha pero tímida.",
        "—La magia curativa es lo único en lo que no me siento tan perdida.",
        "—¿Y nunca has pensado en aprender algún hechizo ofensivo? —le pregunté con curiosidad, apoyando los codos sobre las rodillas—. Una bala de agua o un corte afilado. Aunque sea para tener con qué responder si alguien te arrincona.",
        "Kaoruko bajó un momento la mirada hacia sus manos, jugueteando con los bordes de la toalla.",
        "—Sé que en algún momento será necesario... sobre todo en este lugar —admitió con voz suave—. Pero la verdad es que no me gusta la idea de lastimar a nadie. Además, siento que soy mucho más útil como soporte que atacando en primera línea. Si alguno de mis compañeros sale herido en una misión o en combate... quiero ser capaz de cerrar sus heridas y mantenerlo a salvo.",
        "La miré de reojo. No había cobardía en sus palabras, sino una convicción sincera y protectora.",
        "—Tiene sentido —asentí, encogiéndome de hombros—. Tú encárgate de curar desde atrás. De quebrar huesos y limpiar el camino ya me encargo yo.",
        "Kaoruko parpadeó sorprendida y luego sus mejillas se encendieron con un rubor cálido.",
        "—¡E-Eso suena muy tranquilizador, pero tampoco te arriesgues tanto! —dijo inflando un poco las mejillas en un reproche adorable.",
        "* * *",
        "Estábamos en medio de esa conversación cuando sentí una sombra proyectarse cerca de nosotros.",
        "Por el sendero de adoquines pasaba una chica de cabello plateado claro, recogido en mechones largos que caían sobre su pecho. Tenía una mirada penetrante de ojos verdes y una expresión distante, casi gélida. Llevaba una pequeña bolsa de tela sujeta entre las manos y, al vernos juntos, aminoró el paso y nos clavó una mirada cargada de recelo y desdén.",
        "Kaoruko la vio al instante y sus ojos se iluminaron.",
        "—¡Ah! ¡Subaru-chan! —exclamó, poniéndose de pie de un salto y sacudiéndose la falda.",
        "La recién llegada se detuvo. Su expresión no cambió demasiado, manteniendo ese ceño fruncido y distante, pero cuando Kaoruko corrió hacia ella con una sonrisa brillante, la tensión en sus hombros pareció ceder un poco.",
        "Me quedé sentado observando la escena.",
        "Kaoruko le empezó a hablar con su habitual calidez y entusiasmo. La chica de pelo plateado apenas decía un par de palabras en voz baja, pero asentía con paciencia a todo lo que Kaoruko le contaba, demostrando que sentía un aprecio genuino por ella.",
        "Sin embargo, de cuando en cuando, sus ojos verdes se desviaban hacia mí, mirándome como si fuera una amenaza andante.",
        "«O tiene un problema serio con los hombres, o mi tamaño la tiene con la guardia alta», pensé para mis adentros.",
        "Poco después, Kaoruko la tomó suavemente del brazo y la guio hacia donde yo estaba.",
        "—¡Seier-kun, mira! Ella es Hoshina Subaru —la presentó Kaoruko con una sonrisa orgullosa—. La conocí en los dormitorios el primer día. Puede parecer un poco seria por fuera, ¡pero te juro que es súper amable!",
        "Subaru apretó la bolsita contra su regazo. Al estar a menos de dos metros de mí, pareció notar de golpe mi complexión: hombros anchos como una puerta, cuello grueso y brazos cargados de masa muscular. Tragó saliva con disimulo y dio un paso instintivo hacia atrás, colocándose medio cuerpo detrás de Kaoruko.",
        "Me puse de pie con calma para no parecer descortés.",
        "—Seier Silva —dije con mi tono grave habitual, haciendo un leve gesto con la cabeza—. Mucho gusto, Hoshina.",
        "—Subaru... solo llámame Subaru —respondió ella con voz seca y pausada, manteniendo los ojos fijos en mí—. Silva... sé quién eres. El rumor del coliseo corrió rápido por los pasillos. Dijeron que mandaste a volar a un noble de un solo golpe.",
        "—Solo se tropezó con mi puño —respondí con una media sonrisa tranquila—. No soy de buscar problemas si no me molestan.",
        "Subaru entrecerró los ojos verdes, escéptica.",
        "—Los hombres siempre dicen eso antes de causar un desastre —murmuró, casi para sí misma.",
        "—¡Vamos, Subaru-chan, no seas injusta con él! —intervino Kaoruko haciendo pucheros y cruzándose de brazos—. Seier-kun es muy considerado. Cuando choqué con él el primer día me ayudó a levantar todos mis libros y hasta me regaló pastelillos deliciosos.",
        "A la mención de los pastelillos, los ojos de Subaru parpadearon y miró la bolsa que sostenía entre las manos, pareciendo dudar.",
        "—¿Pastelillos...? —murmuró con una pizca de curiosidad que no logró disimular del todo.",
        "—¡Sí! De miel y frutos rojos —asintió Kaoruko con entusiasmo—. Y de hecho, todavía nos quedan un par en la caja. ¿Quieres uno?",
        "Subaru miró a Kaoruko, luego me miró a mí como evaluando si había algún truco oculto, y finalmente suspiró dejando caer levemente los hombros.",
        "—...Si insistes, supongo que puedo aceptar —dijo en un tono que intentaba sonar indiferente, aunque un tenue tinte rosado asomó en la punta de sus orejas.",
        "Sonreí de lado mientras me agachaba a buscar la caja en mi mochila.",
        "Parecía que el grupo de la Clase F-2 se estaba volviendo bastante más animado de lo previsto."
    ]
}

def compilar_capitulo_4():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4]
    
    # 1. Guardar texto plano del Capítulo 4
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "04_Capitulo_4.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_4['header'].upper()}: {CAPITULO_4['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_4['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 4 con ilustración
    docx_cap4 = os.path.join(NOVELA_DIR, "capitulos", "04_Capitulo_4.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_4['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_4['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)
    
    if os.path.exists(CAPITULO_4['illustration']):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(15)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        r_img.add_picture(CAPITULO_4['illustration'], width=Inches(3.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(20)
        r_cap = p_cap.add_run(CAPITULO_4['illustration_caption'])
        r_cap.font.name = "Georgia"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(120, 100, 140)

    for p_text in CAPITULO_4['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap4)
    print(f"[OK] Word individual Capítulo 4 guardado en: {docx_cap4}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 + Cap 2 + Cap 3 + Cap 4)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_4()
