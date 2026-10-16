import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_novel_document(title, subtitle, author, chapters, output_path):
    doc = Document()
    
    # Page setup - Margins (1 inch / 2.54 cm standard editorial format)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Header (empty for clean look or can have novel title)
        header = section.header
        h_p = header.paragraphs[0]
        h_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = h_p.add_run(f"{title} — {author}")
        h_run.font.name = "Georgia"
        h_run.font.size = Pt(8.5)
        h_run.font.italic = True
        h_run.font.color.rgb = RGBColor(150, 150, 160)
        
        # Footer with page number
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run("— ")
        f_run.font.name = "Georgia"
        f_run.font.size = Pt(9.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)
        
        # XML Field for dynamic page numbers in Word
        f_xml = f_p.add_run()
        f_xml.font.name = "Georgia"
        f_xml.font.size = Pt(9.5)
        f_xml.font.color.rgb = RGBColor(120, 120, 120)
        f_pld = OxmlElement('w:fldSimple')
        f_pld.set(qn('w:instr'), 'PAGE')
        f_xml._r.append(f_pld)
        
        f_run_end = f_p.add_run(" —")
        f_run_end.font.name = "Georgia"
        f_run_end.font.size = Pt(9.5)
        f_run_end.font.color.rgb = RGBColor(120, 120, 120)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Georgia'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 35)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(4)

    # ================= PORTADA =================
    cover_p_space = doc.add_paragraph()
    cover_p_space.paragraph_format.space_before = Pt(100)
    
    # Título Principal
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title.upper())
    r_title.font.name = 'Georgia'
    r_title.font.size = Pt(30)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(35, 25, 50)
    
    # Ornamento
    p_ornament = doc.add_paragraph()
    p_ornament.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ornament.paragraph_format.space_before = Pt(8)
    p_ornament.paragraph_format.space_after = Pt(8)
    r_ornament = p_ornament.add_run("❖   ❖   ❖")
    r_ornament.font.name = 'Georgia'
    r_ornament.font.size = Pt(14)
    r_ornament.font.color.rgb = RGBColor(140, 95, 160)
    
    # Subtítulo
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(20)
        r_sub = p_sub.add_run(subtitle)
        r_sub.font.name = 'Georgia'
        r_sub.font.size = Pt(14)
        r_sub.font.italic = True
        r_sub.font.color.rgb = RGBColor(85, 75, 105)
    
    p_author_space = doc.add_paragraph()
    p_author_space.paragraph_format.space_before = Pt(140)
    
    # Autor
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_by = p_author.add_run("Una historia de\n")
    r_by.font.name = 'Georgia'
    r_by.font.size = Pt(11)
    r_by.font.color.rgb = RGBColor(130, 130, 140)
    
    r_author = p_author.add_run(author)
    r_author.font.name = 'Georgia'
    r_author.font.size = Pt(17)
    r_author.font.bold = True
    r_author.font.color.rgb = RGBColor(45, 40, 60)

    # Banner de advertencia +18
    p_notice = doc.add_paragraph()
    p_notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_notice.paragraph_format.space_before = Pt(45)
    r_notice = p_notice.add_run("✦ ADVERTENCIA: NOVELA PARA ADULTOS (+18) ✦\nEsta obra contiene escenas de combate visceral, lenguaje maduro, romance explícito\ny contenido sensual no apto para menores de edad.")
    r_notice.font.name = 'Georgia'
    r_notice.font.size = Pt(9.5)
    r_notice.font.italic = True
    r_notice.font.color.rgb = RGBColor(165, 45, 55)

    doc.add_page_break()

    # ================= HOJA DE FICHA / SINOPSIS =================
    p_prof_head = doc.add_paragraph()
    p_prof_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prof_head.paragraph_format.space_before = Pt(25)
    p_prof_head.paragraph_format.space_after = Pt(15)
    r_pf = p_prof_head.add_run("REGISTRO DEL PROTAGONISTA")
    r_pf.font.name = 'Georgia'
    r_pf.font.size = Pt(13)
    r_pf.font.bold = True
    r_pf.font.color.rgb = RGBColor(90, 60, 120)

    # Tabla de características
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False

    profile_data = [
        ("Nombre:", "Seier Silva"),
        ("Edad y Físico:", "16 años | 1,75 m | 90 kg de masa muscular densa y compacta"),
        ("Origen:", "Reencarnado de la Tierra | Hijo único de la caída Casa Baronal Silva"),
        ("Estilo Marcial:", "Híbrido de Muay Thai y Jeet Kune Do (combate a corta distancia, clinch, codos, rodillas y desvíos directos)"),
        ("Afinidad / Energía:", "Canalización física de maná interior (reforzamiento neuromuscular y masa de impacto)"),
        ("Personalidad:", "Tosco al inicio, pragmático, observador, de pocas palabras pero leal y protector una vez que entra en confianza. Constante sed de aprendizaje.")
    ]

    for i, (label, val) in enumerate(profile_data):
        row = table.rows[i]
        
        # Columna 1
        p_l = row.cells[0].paragraphs[0]
        p_l.paragraph_format.space_after = Pt(3)
        r_l = p_l.add_run(label)
        r_l.font.name = 'Georgia'
        r_l.font.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = RGBColor(50, 45, 60)
        row.cells[0].width = Inches(1.8)
        
        # Columna 2
        p_v = row.cells[1].paragraphs[0]
        p_v.paragraph_format.space_after = Pt(3)
        r_v = p_v.add_run(val)
        r_v.font.name = 'Georgia'
        r_v.font.size = Pt(10)
        r_v.font.color.rgb = RGBColor(40, 40, 45)
        row.cells[1].width = Inches(4.5)

    p_post_table = doc.add_paragraph()
    p_post_table.paragraph_format.space_before = Pt(30)
    p_post_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div_prof = p_post_table.add_run("❖   ❖   ❖")
    r_div_prof.font.name = 'Georgia'
    r_div_prof.font.size = Pt(12)
    r_div_prof.font.color.rgb = RGBColor(160, 140, 180)

    doc.add_page_break()

    # ================= CAPÍTULOS =================
    for chap in chapters:
        # Título del Capítulo
        p_chap = doc.add_paragraph()
        p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_chap.paragraph_format.space_before = Pt(35)
        p_chap.paragraph_format.space_after = Pt(4)
        
        r_num = p_chap.add_run(chap.get('header', '').upper() + "\n")
        r_num.font.name = 'Georgia'
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(140, 85, 160)
        
        r_name = p_chap.add_run(chap.get('title', ''))
        r_name.font.name = 'Georgia'
        r_name.font.size = Pt(20)
        r_name.font.bold = True
        r_name.font.color.rgb = RGBColor(30, 25, 45)
        
        # Divisor decorativo
        p_div = doc.add_paragraph()
        p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_div.paragraph_format.space_after = Pt(26)
        r_div = p_div.add_run("— ✦ —")
        r_div.font.name = 'Georgia'
        r_div.font.size = Pt(11)
        r_div.font.color.rgb = RGBColor(160, 145, 180)
        
        # Ilustraciones del capitulo si existen
        imgs_to_add = []
        if chap.get('illustrations'):
            imgs_to_add.extend(chap['illustrations'])
        elif chap.get('illustration') and os.path.exists(chap['illustration']):
            imgs_to_add.append({
                'image': chap['illustration'],
                'caption': chap.get('illustration_caption', 'Ilustración del Capítulo')
            })
            
        for item in imgs_to_add:
            img_path = item.get('image') if isinstance(item, dict) else item
            cap_text = item.get('caption', 'Ilustración') if isinstance(item, dict) else 'Ilustración'
            if img_path and os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(3.2))
                
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(16)
                r_cap = p_cap.add_run(cap_text)
                r_cap.font.name = 'Georgia'
                r_cap.font.size = Pt(9.5)
                r_cap.font.italic = True
                r_cap.font.color.rgb = RGBColor(130, 110, 150)
        
        # Párrafos
        for paragraph_text in chap.get('paragraphs', []):
            paragraph_text = paragraph_text.strip()
            if not paragraph_text:
                continue
            
            if paragraph_text.startswith("### "):
                # Subtítulo de sección interna
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.space_before = Pt(18)
                p_sub.paragraph_format.space_after = Pt(6)
                r_sub = p_sub.add_run(paragraph_text[4:])
                r_sub.font.name = 'Georgia'
                r_sub.font.size = Pt(12)
                r_sub.font.bold = True
                r_sub.font.color.rgb = RGBColor(85, 45, 120)
            elif paragraph_text in ["* * *", "✦ ✦ ✦", "❖ ❖ ❖"]:
                # Separador de escena
                p_sep = doc.add_paragraph()
                p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_sep.paragraph_format.space_before = Pt(14)
                p_sep.paragraph_format.space_after = Pt(14)
                r_sep = p_sep.add_run("✦   ✦   ✦")
                r_sep.font.name = 'Georgia'
                r_sep.font.size = Pt(11)
                r_sep.font.color.rgb = RGBColor(160, 130, 180)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.25
                p.paragraph_format.space_after = Pt(5)
                
                # Formateo de sangrías y diálogos
                if paragraph_text.startswith("—") or paragraph_text.startswith("- "):
                    # Diálogo narrativo en español
                    p.paragraph_format.left_indent = Inches(0.18)
                else:
                    # Párrafo descriptivo estándar de novela
                    p.paragraph_format.first_line_indent = Inches(0.28)
                
                # Añadir texto
                r = p.add_run(paragraph_text)
                r.font.name = 'Georgia'
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(28, 28, 32)

        doc.add_page_break()

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"[OK] Novela compilada exitosamente en:\n{output_path}")

def create_single_chapter_document(chap, output_path, novel_title="El camino del reencarnado", author="Seier Silva"):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Encabezado
        header = section.header
        h_p = header.paragraphs[0]
        h_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = h_p.add_run(f"{novel_title} — {chap.get('header', '')}")
        h_run.font.name = "Georgia"
        h_run.font.size = Pt(8.5)
        h_run.font.italic = True
        h_run.font.color.rgb = RGBColor(150, 150, 160)
        
        # Pie de página con número de página dinámico
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = f_p.add_run("— ")
        f_run.font.name = "Georgia"
        f_run.font.size = Pt(9.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)
        f_xml = f_p.add_run()
        f_xml.font.name = "Georgia"
        f_xml.font.size = Pt(9.5)
        f_xml.font.color.rgb = RGBColor(120, 120, 120)
        f_pld = OxmlElement('w:fldSimple')
        f_pld.set(qn('w:instr'), 'PAGE')
        f_xml._r.append(f_pld)
        f_run_end = f_p.add_run(" —")
        f_run_end.font.name = "Georgia"
        f_run_end.font.size = Pt(9.5)
        f_run_end.font.color.rgb = RGBColor(120, 120, 120)

    # Estilo Normal
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Georgia'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 35)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(4)

    # Título del Capítulo
    p_chap = doc.add_paragraph()
    p_chap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_chap.paragraph_format.space_before = Pt(30)
    p_chap.paragraph_format.space_after = Pt(4)
    
    r_num = p_chap.add_run(chap.get('header', '').upper() + "\n")
    r_num.font.name = 'Georgia'
    r_num.font.size = Pt(12)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(140, 85, 160)
    
    r_name = p_chap.add_run(chap.get('title', ''))
    r_name.font.name = 'Georgia'
    r_name.font.size = Pt(20)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(30, 25, 45)
    
    # Divisor decorativo
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(24)
    r_div = p_div.add_run("— ✦ —")
    r_div.font.name = 'Georgia'
    r_div.font.size = Pt(11)
    r_div.font.color.rgb = RGBColor(160, 145, 180)

    # Ilustraciones del capitulo si existen
    imgs_to_add = []
    if chap.get('illustrations'):
        imgs_to_add.extend(chap['illustrations'])
    elif chap.get('illustration') and os.path.exists(chap['illustration']):
        imgs_to_add.append({
            'image': chap['illustration'],
            'caption': chap.get('illustration_caption', 'Ilustración del Capítulo')
        })
        
    for item in imgs_to_add:
        img_path = item.get('image') if isinstance(item, dict) else item
        cap_text = item.get('caption', 'Ilustración') if isinstance(item, dict) else 'Ilustración'
        if img_path and os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            r_img = p_img.add_run()
            r_img.add_picture(img_path, width=Inches(3.2))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(16)
            r_cap = p_cap.add_run(cap_text)
            r_cap.font.name = 'Georgia'
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(130, 110, 150)

    # Párrafos
    for paragraph_text in chap.get('paragraphs', []):
        paragraph_text = paragraph_text.strip()
        if not paragraph_text:
            continue
        
        if paragraph_text.startswith("### "):
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(18)
            p_sub.paragraph_format.space_after = Pt(6)
            r_sub = p_sub.add_run(paragraph_text[4:])
            r_sub.font.name = 'Georgia'
            r_sub.font.size = Pt(12)
            r_sub.font.bold = True
            r_sub.font.color.rgb = RGBColor(85, 45, 120)
        elif paragraph_text in ["* * *", "✦ ✦ ✦", "❖ ❖ ❖"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(14)
            p_sep.paragraph_format.space_after = Pt(14)
            r_sep = p_sep.add_run("✦   ✦   ✦")
            r_sep.font.name = 'Georgia'
            r_sep.font.size = Pt(11)
            r_sep.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if paragraph_text.startswith("—") or paragraph_text.startswith("- "):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            
            r = p.add_run(paragraph_text)
            r.font.name = 'Georgia'
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"[OK] Capítulo guardado en Word en:\n{output_path}")


