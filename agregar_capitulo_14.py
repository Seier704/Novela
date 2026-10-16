# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_14 = {
    "header": "Capítulo 14",
    "title": "Dúo de Fuego y Travesuras",
    "paragraphs": [
        "El último día de descanso antes de retomar la rutina habitual de las clases fue, sin lugar a dudas, el más caótico de la semana.",
        "Y la culpa absoluta la tuvieron dos personas: Ken y Burnice White.",
        "Desde que regresamos del bosque, esos dos habían descubierto que compartían exactamente la misma neurona hiperactiva y la misma devoción desmedida por las llamas. Ken vivía para la adrenalina y el calor de las artes marciales; Burnice vivía para mezclar combustibles, hacer estallar cosas y sonreír como si el mundo entero fuera un festival de fuegos artificiales.",
        "Cuando dos espíritus libres como ellos se juntaban, la tranquilidad en los dormitorios de la Clase F quedaba cancelada de inmediato.",
        "* * *",
        "La jornada comenzó a media mañana en el patio de los lavaderos.",
        "Estaba sentado en un banco afilando con calma las hebillas de mi abrigo de cuero junto a Kaoruko, que repasaba un libro botánico sobre ungüentos, cuando un estallido sordo seguido de una nube de humo con un sospechoso aroma a azúcar quemada y queroseno se elevó cerca de las cuerdas de tender.",
        "—¡Waaaaah! ¡Se está quemando la sábana! —chilló Ken a todo pulmón.",
        "—¡Tranquilo, tranquilo, tengo el extintor de soda! —se escuchó la voz eufórica de Burnice.",
        "Un segundo después, una llamarada rosada y chispeante devoró una esquina de la ropa tendida justo antes de que un chorro de espuma gaseosa la apagara con un siseo estruendoso.",
        "De la nube de humo salieron corriendo los dos, con las caras llenas de hollín y sosteniendo sendos palos de madera con salchichas carbonizadas en las puntas.",
        "—¡LES DIJE QUE NO ASARAN COMIDA CON EL COMBUSTIBLE NITRO EN MI LAVANDERÍA, PEDAZOS DE ANIMALES! —rugió Lucy desde el umbral, con su gorra militar torcida y blandiendo su garrote con una furia cómica.",
        "—¡Corre, Burnice, que la capitana pequeña viene con ganas de sangre! —gritó Ken riendo a carcajadas mientras saltaba la valla del huerto con una voltereta limpia.",
        "—¡Jajajaja, pero si el glaseado de miel quedó crujiente, Lucy! —respondió Burnice dando brincos con sus gafas rojas temblando en su frente.",
        "Lucy los persiguió por todo el jardín arrojándoles piedras que ambos esquivaban entre piruetas y carcajadas, haciendo que Kaoruko se tapara la boca con las manos para ahogar la risa y Subaru Hoshina suspirara desde la sombra de los cerezos.",
        "—Esos dos tienen un tornillo suelto —murmuró Subaru Hoshina, aunque una pequeña sonrisa delataba que se estaba divirtiendo tanto como nosotros.",
        "Poco después, la pareja de pirómanos fue a molestar a Piper Wheel, que dormía plácidamente acurrucada sobre el banco de herramientas. Burnice le colocó unas gafas de sol con forma de corazón en la nariz mientras Ken le tarareaba una tonada marcial al oído; Piper simplemente abrió un ojo verde somnoliento, bostezó como un gato perezoso y les arrojó una pequeña tuerca que rebotó en la frente de Ken antes de volver a dormirse.",
        "* * *",
        "Para el mediodía, agotadas sus travesuras por los dormitorios, Ken y Burnice decidieron que necesitaban gastar energía de verdad.",
        "—¡Oye, Seier! ¡Kaoruko! —llamó Ken acercándose a las gradas de la arena de tierra, pasándose un brazo por la frente para secarse el sudor—. ¡Vamos a hacer un sparring entre Burnice y yo! Tienen que ver esto.",
        "—¡Sí, vengan! ¡Prometo no encender fuego en las gradas... creo! —añadió Burnice con un guiño alegre.",
        "Varios de nosotros nos acomodamos en la barandilla de madera para ver el espectáculo: Kaoruko, Subaru Hoshina, Lighter con sus gafas oscuras y yo.",
        "Ken y Burnice se plantaron en el centro de la arena.",
        "A diferencia de los duelos tensos contra nobles o bestias, esto era pura diversión entre camaradas. Ken se ajustó la bandana roja y levantó su guardia de artes marciales con los puños chisporroteando en llamas doradas. Burnice ajustó las boquillas de sus propulsores manuales, encendiendo dos lengüetas de fuego carmesí con un chasquido metálico.",
        "—¡Sin llorar si te quemo el flequillo azul, Ken-kun! —provocó ella sonriendo de oreja a oreja.",
        "—¡Atrévete si puedes alcanzarme, reina del caos!",
        "El combate fue un torbellino de luces, acrobacias y ritmo frenético.",
        "Burnice se movía con un estilo impredecible, deslizándose por la arena con impulsos cortos de fuego en sus botas, girando como una peonza mientras lanzaba arcos de chispas ardientes que iluminaban toda la pista. Ken respondía con fintas aéreas, esquivando las llamaradas por milímetros y contraatacando con patadas circulares envueltas en llamas que disipaban el combustible en el aire.",
        "¡FWOOSH! ¡CRACK!",
        "En un momento, Burnice lanzó una cortina de humo dulce que cegó temporalmente a Ken, intentando sorprenderlo por la espalda con un golpe acolchado, pero Ken reaccionó por instinto: giró sobre su talón, se agachó y le barrió los pies con una barrida baja.",
        "Burnice cayó de espaldas sobre la arena, pero antes de tocar el suelo activó una pequeña explosión en su palma libre que la impulsó hacia arriba de nuevo, aterrizando encima del pecho de Ken y haciéndolos rodar a los dos por la tierra entre risas descontroladas.",
        "—¡Punto para mí por el derribo! —gritó Ken desde el suelo.",
        "—¡Punto para mí por la explosión en tu barbilla! —contraatacó Burnice despeinándole el cabello.",
        "* * *",
        "Ambos terminaron tumbados bocarriba en la hierba del borde de la arena, jadeando por el esfuerzo pero riendo a carcajadas con los brazos abiertos hacia el cielo despejado.",
        "—Ustedes dos son un peligro público —dijo Lighter apoyado en la barandilla, aunque sonreía con orgullo al ver la coordinación y velocidad de ambos.",
        "—¡Pero fue divertidísimo! —exclamó Burnice, sentándose y sacudiéndose el polvo de la falda.",
        "Kaoruko bajó las gradas con una jarra de agua fresca y un par de paños limpios, entregándoselos con esa ternura que siempre la caracterizaba.",
        "—Tomen un poco de agua fresca... miren cómo quedaron, llenos de ceniza —les dijo con una sonrisa dulce.",
        "—¡Gracias, Kaoruko, eres un ángel! —dijo Ken bebiendo de un solo trago.",
        "Me acerqué a ellos y les choqué los puños en señal de respeto.",
        "—Buen intercambio de llamas —les dije con una sonrisa tranquila—. Tienen buena química peleando juntos.",
        "—¡Obvio! ¡Cuando entremos a una batalla grande, seremos el dúo incinerador de la F-2! —aseguró Burnice con un brillo travieso en la mirada.",
        "El sol de la tarde comenzó a descender, marcando el final de los días libres.",
        "Al ver a Ken y Burnice bromear, a Kaoruko sonreír y a la clase entera convivir en armonía, el aire de la academia ya no se sentía frío ni ajeno. Mañana volverían las clases formales y las miradas hostiles de los nobles, pero nuestra manada estaba más unida y encendida que nunca."
    ]
}

def compilar_capitulo_14():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    from agregar_capitulo_12 import CAPITULO_12
    from agregar_capitulo_13 import CAPITULO_13
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12, CAPITULO_13, CAPITULO_14]
    
    # 1. Guardar texto plano del Capítulo 14
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "14_Capitulo_14.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_14['header'].upper()}: {CAPITULO_14['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_14['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 14
    docx_cap14 = os.path.join(NOVELA_DIR, "capitulos", "14_Capitulo_14.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_14['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_14['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_14['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap14)
    print(f"[OK] Word individual Capítulo 14 guardado en: {docx_cap14}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 14)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_14()
