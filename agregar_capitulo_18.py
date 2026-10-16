# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_18 = {
    "header": "Capítulo 18",
    "title": "Vapor, Tinta y Piel",
    "paragraphs": [
        "Avanzada la medianoche, el bullicio del banquete fue apagándose poco a poco.",
        "El cansancio de semanas de esfuerzo y el efecto relajante de los baños hicieron mella en la mayoría: Piper ya dormía abrazada a un cojín, Yami se había retirado a fumar bajo los aleros y Kaoruko bostezaba con una timidez adorable antes de marcharse a los dormitorios acompañada por Subaru Hoshina.",
        "Sin embargo, antes de retirarnos, una pequeña conspiración se había gestado en la mesa.",
        "Era un secreto a voces que entre Renji y Rukia existía algo mucho más profundo que una simple amistad de armas. Crecieron juntos en los barrios más duros de la provincia sur, cuidándose las espaldas desde niños descalzos en el barro; pero por el orgullo tosco de Renji y la reserva aristocrática de Rukia, jamás se habían atrevido a dar el primer paso. Se trataban a gritos y desplantes para ocultar lo que ardía debajo.",
        "—Oigan —había susurrado Natsuki mientras recogíamos los platos, guiñándonos un ojo a Ken, a Lucy y a mí—. Dejé la llave del pabellón de aguas termales exteriores 'olvidada' en la mesa de Renji... y le dije a Rukia que fuera a buscar su toalla favorita al borde del estanque.",
        "—Si nos descubren, Erza nos va a decapitar —bromeó Ken entre risas bajas.",
        "—No van a descubrir nada —dije yo con una media sonrisa—. Esos dos necesitaban este empujón hace años.",
        "Apagamos las linternas comunes y dejamos que la noche hiciera el resto.",
        "* * *",
        "El estanque exterior de aguas termales estaba completamente desierto bajo el manto de la noche.",
        "Una niebla densa de vapor blanco danzaba sobre la superficie del agua mineral, iluminada únicamente por la luz de la luna llena y las linternas de piedra cubiertas de musgo. El silencio solo era roto por el murmullo continuo de la cascada tibia cayendo entre los peñascos.",
        "Rukia llegó envuelta en un yukata azul oscuro, mirando a ambos lados con el ceño fruncido.",
        "—¿Dónde dejó la toalla ese idiota de Natsuki...? —murmuró para sí misma.",
        "—No hay ninguna toalla, Rukia.",
        "La voz profunda y algo ronca de Renji la hizo sobresaltarse. El pelirrojo estaba sentado en la orilla sobre una roca plana, descalzo, con su yukata entreabierto en el pecho dejando ver los tatuajes tribales oscuros que serpenteaban por su clavícula y su torso.",
        "—¡Renji! ¿Qué haces aquí a estas horas? —inquirió ella, cruzándose de brazos para disimular el súbito vuelco que sintió en el pecho—. Si nos ven solos en las aguas termales mixtas a medianoche, los profesores...",
        "—Al cuerno los profesores —la cortó Renji, poniéndose de pie con lentitud. Su altura imponente de guerrero contrastaba con la figura menuda y grácil de ella—. Los chicos planearon esto. Y la verdad... me alegro de que lo hicieran.",
        "Rukia dio un paso involuntario hacia atrás, pero sus talones chocaron con el borde húmedo del muelle.",
        "—¿De qué estás hablando?",
        "—De ti y de mí —dijo Renji, acortando la distancia con pasos firmes hasta quedar a escasos centímetros de su rostro. Sus ojos marrones, siempre feroces en la batalla, desbordaban ahora una vulnerabilidad y una ternura que la desarmaron por completo—. Llevamos años jugando a pelearnos, fingiendo que solo somos compañeros de armas. Pero desde que éramos unos mocosos muertos de hambre en las calles... lo único que me ha importado en este mundo eres tú.",
        "A Rukia se le cortó la respiración. Sus grandes ojos violetas se abrieron con asombro, reflejando el brillo de la luna.",
        "—Renji...",
        "—Siempre tuve miedo de decírtelo —confesó él en un susurro áspero, levantando una mano ancha y callosa para acariciar con extrema delicadeza la mejilla fría de la chica—. Miedo de que fueras demasiado para mí... de que te alejaran de mi lado. Pero en el calabo esta semana, cuando vi a esos monstruos caer... supe que no puedo seguir guardándome esto. Te amo, Rukia. Te he amado desde el primer día.",
        "Una lágrima solitaria rodó por la mejilla de Rukia, perdiéndose bajo los dedos de Renji. La barrera de compostura y orgullo que había sostenido durante tanto tiempo se derrumbó en un segundo.",
        "—Eres un idiota, Renji... —sollozó ella con una sonrisa temblorosa, apoyando su frente contra el pecho desnudo del muchacho—. Un completo idiota por haber tardado tanto en decírmelo.",
        "Renji no esperó más. Inclinó el rostro y atrapó sus labios en un beso.",
        "Fue un beso desesperado y hambriento al principio, cargado con años de anhelo contenido, pero que pronto se volvió profundo, cálido y devorador. Rukia enredó sus pequeños dedos en la nuca de Renji, aferrándose a su cabello escarlata mientras sus bocas se abrían y sus lenguas se encontraban en una danza ardiente que hizo que ambos se estremecieran.",
        "* * *",
        "Sin romper el contacto de sus labios, las manos de Renji se deslizaron hacia el lazo del yukata de Rukia, desatándolo con dedos temblorosos. La tela oscura cayó suavemente sobre las tablas de madera, revelando la silueta nívea y perfecta de la chica bajo la luz de la luna: senos pequeños y firmes de pezones sonrosados que se erizaron con el aire fresco, una cintura estrecha y caderas suaves y proporcionadas.",
        "Renji tragó saliva, contemplándola como si tuviera a una diosa frente a él.",
        "—Eres hermosa... increíblemente hermosa, Rukia —susurró con veneración.",
        "Rukia se sonrojó hasta la raíz del cabello, pero no apartó la mirada. Sus manos delgadas tiraron del yukata de Renji hacia abajo, dejando al descubierto su físico imponente, cubierto de cicatrices de batalla y tatuajes que resaltaban cada músculo tenso de su espalda y abdomen.",
        "Tomados de la mano, descendieron juntos los escalones de piedra hacia el agua termal.",
        "El agua caliente los envolvió como un manto tibio y líquido. Rukia soltó un suspiro entrecortado cuando el calor acarició su piel desnuda, pero el calor que más la quemaba era el del cuerpo de Renji pegado al suyo.",
        "Renji la atrajo hacia sí en el centro del estanque, donde el agua les llegaba al pecho. La acorraló suavemente contra la pared lisa de roca volcánica. Sus manos grandes y firmes se posaron en la cintura de Rukia, sintiendo la suavidad sedosa de su piel húmeda, mientras sus labios descendían por su mandíbula, besando la curva de su cuello y haciéndola arquear la espalda con un gemido suave.",
        "—Renji... ah... —jadeó Rukia, cerrando los ojos mientras los dedos de él subían para acariciar sus pechos, masajeándolos con devoción y rozando sus pezones endurecidos bajo el agua.",
        "—Dime si quieres que pare —murmuró él con la respiración entrecortada contra su oído.",
        "—No te atrevas a parar... hazme tuya, Renji —le pidió ella, mirándolo con una pasión encendida que él jamás había visto en sus ojos.",
        "Renji deslizó una mano por debajo del agua, acariciando el vientre plano de Rukia y descendiendo entre sus muslos húmedos. Sus dedos encontraron la intimidad femenina, empapada no solo por las aguas termales sino por el deseo desbordante que la consumía. Al acariciar su punto más sensible, Rukia mordió el hombro de Renji para sofocar un grito de placer.",
        "Con cuidado y fuerza contenida, Renji levantó a Rukia por los muslos, sentándola en el borde de un saliente de roca sumergido. Ella abrió las piernas alrededor de sus caderas, sintiendo la dureza ardiente y pulsante de la virilidad de él presionando contra su entrada.",
        "Se miraron a los ojos durante un latido eterno. No había dudas, solo una certeza absoluta.",
        "Renji empujó hacia adelante despacio, hundiéndose en ella.",
        "Rukia soltó un jadeo profundo y sonoro cuando la plenitud de Renji la llenó por completo. La sensación de invasión y entrega era tan intensa que sus uñas se clavaron en los hombros tatuados de él, dejando marcas rojas en su piel. Renji contuvo la respiración, embriagado por la estrechez ardiente y húmeda de Rukia abrazándolo por dentro.",
        "—Rukia... estás tan cálida... tan perfecta —gimió Renji con voz quebrada.",
        "—Se siente... tan bien... muévete, Renji... por favor...",
        "Renji comenzó a moverse con embestidas lentas y profundas, cuidando cada roce antes de que el ritmo se acelerara inevitablemente. El agua del estanque chapoteaba rítmicamente contra las rocas al compás de sus cuerpos uniéndose una y otra vez. Rukia echó la cabeza hacia atrás sobre la piedra, con el cabello negro empapado pegado a su cuello, dejando escapar gemidos húmedos y entrecortados que el murmullo de la cascada amortiguaba en la noche.",
        "—¡Renji... ah... más... así...! —suplicaba ella, envolviendo con más fuerza sus piernas alrededor de su cintura, empujando sus caderas contra las de él para recibirlo hasta el fondo.",
        "Renji perdió todo vestigio de control. Sostuvo las caderas de Rukia con un agarre de hierro y aumentó la fuerza y la velocidad de sus embestidas, hundiéndose en ella con una pasión salvaje que hacía que el vapor a su alrededor pareciera arder. Cada golpe dentro de ella resonaba en el alma de ambos, borrando años de soledad, de peleas fingidas y de dolor.",
        "La tensión acumulada en el vientre de Rukia explotó de pronto en una ola devastadora de placer.",
        "—¡REEEENJI! —gritó en un susurro ahogado, arqueando el torso mientras su interior se contraía con espasmos violentos y deliciosos alrededor del miembro de él, derramándose en un clímax abrumador.",
        "Sentir las contracciones de Rukia fue el detonante final para Renji. Con un gruñido gutural y profundo que le desgarró la garganta, dio dos últimas embestidas a fondo y se derramó dentro de ella con una descarga ardiente y prolongada, fundiéndose en un abrazo desesperado mientras ambos temblaban en el éxtasis compartido.",
        "* * *",
        "Permanecieron abrazados en el agua durante un largo tiempo, con las frentes unidas y las respiraciones calmándose al compás de la cascada.",
        "Renji la cubrió con sus brazos protectores, besando tiernamente su frente, sus párpados y sus labios exhaustos. Rukia sonreía acurrucada contra su pecho, sintiendo el latido desbocado del corazón de Renji contra su mejilla.",
        "Antes del amanecer, recogieron sus ropas en silencio y regresaron a sus respectivas alas de la suite con pasos ligeros y cómplices, asegurándose de que nadie los viera.",
        "Al día siguiente, durante el desayuno matutino, ambos se sentaron a la mesa actuando como siempre: Renji comiendo carne con las manos y Rukia regañándolo por sus modales campesinos.",
        "Pero cuando sus miradas se cruzaron por encima de las tazas de té, un brillo secreto, ardiente y eterno delató que aquella noche en las aguas termales lo había cambiado todo para siempre."
    ]
}

def compilar_capitulo_18():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    from agregar_capitulo_12 import CAPITULO_12
    from agregar_capitulo_13 import CAPITULO_13
    from agregar_capitulo_14 import CAPITULO_14
    from agregar_capitulo_15 import CAPITULO_15
    from agregar_capitulo_16 import CAPITULO_16
    from agregar_capitulo_17 import CAPITULO_17
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12, CAPITULO_13, CAPITULO_14, CAPITULO_15, CAPITULO_16, CAPITULO_17, CAPITULO_18]
    
    # 1. Guardar texto plano del Capítulo 18
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "18_Capitulo_18.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_18['header'].upper()}: {CAPITULO_18['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_18['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 18
    docx_cap18 = os.path.join(NOVELA_DIR, "capitulos", "18_Capitulo_18.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_18['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_18['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_18['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap18)
    print(f"[OK] Word individual Capítulo 18 guardado en: {docx_cap18}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 18)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_18()
