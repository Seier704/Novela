# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_13 = {
    "header": "Capítulo 13",
    "title": "Sombras y Determinación",
    "paragraphs": [
        "El segundo día de descanso concedido por la academia me levanté temprano para trotar por el perímetro de los establos.",
        "El aire de la mañana era frío y limpio. Creí que sería el único despierto a esa hora, pero al pasar cerca del patio de armas este, el silbido torpe de una espada de madera cortando el aire llamó mi atención.",
        "Me acerqué despacio y me apoyé contra la valla de madera.",
        "En el centro de la arena estaba Subaru Natsuki. Llevaba su chándal habitual con las mangas remangadas hasta los codos, el rostro bañado en sudor y las manos vendadas sosteniendo una espada de práctica de fresno. Repetía una y otra vez el mismo corte descendente contra un muñeco de paja, jadeando como si llevara horas allí.",
        "Su técnica era desordenada y carecía de fuerza en las caderas, pero la terquedad con la que volvía a levantar los brazos delataba algo más que una simple curiosidad.",
        "* * *",
        "—Si no afirmas la pierna de apoyo, la vibración del impacto te va a abrir las muñecas —le dije con mi voz tranquila, saltando la barandilla.",
        "Natsuki dio un brinco del susto, casi soltando el madero.",
        "—¡A-Ah! ¡Seier! —exclamó, secándose el sudor de los ojos con el dorso del brazo mientras respiraba con dificultad—. Casi me provocas un infarto, hombre de granito. ¿Cuánto tiempo llevas ahí mirando?",
        "—El suficiente para ver que tienes los hombros acalambrados —respondí, acercándome y apoyando un pie en una piedra—. No pensé que te interesara la esgrima.",
        "Natsuki bajó la espada de madera y la clavó en la tierra, mirando sus propias manos enrojecidas con una mezcla de frustración y vergüenza.",
        "—Bueno... en la expedición del bosque me di cuenta de muchas cosas —admitió en voz baja, perdiendo su habitual tono bromista—. Soy bueno planeando mapas, cargando provisiones y gritando advertencias desde la retaguardia. Pero cuando ese lobo o ese gólem se nos vinieron encima... no pude hacer absolutamente nada. Si tú, Caesar o Erza no hubieran estado ahí... yo habría sido el primero en morir.",
        "—Hiciste tu parte, Natsuki. La logística salvó al equipo.",
        "—¡Pero no quiero ser solo una carga cuando las cosas se pongan feas de verdad! —replicó él, apretando los puños con fuerza—. Odio ser el tipo débil al que todos tienen que proteger. Quiero aprender a blandir un filo... y tal vez, solo tal vez, hacer funcionar mi magia.",
        "—¿Tu magia? —arqueé una ceja—. Nunca te he visto canalizar en las clases.",
        "Natsuki titubeó un momento, miró hacia ambos lados como si fuera a revelar un secreto prohibido, y luego extendió la palma abierta hacia el suelo.",
        "Cerró los ojos y apretó los dientes. Tras unos segundos de esfuerzo visible, un tenue velo oscuro, denso y nebuloso como humo de alquitrán, comenzó a brotar de sus dedos, devorando la luz del sol alrededor de su mano antes de disiparse en una ráfaga inestable.",
        "—Elemento Oscuridad —dijo Natsuki con una mueca amarga—. En la capital dicen que es una afinidad maldita, difícil de controlar y sin los hechizos vistosos del fuego o el rayo. Nadie me ha enseñado cómo usarla sin que me consuma el maná en cinco segundos.",
        "* * *",
        "—Menuda queja de llorón para empezar la mañana.",
        "Una voz profunda, áspera y cargada de tabaco retumbó detrás de nosotros.",
        "Natsuki se congeló. Ambos nos giramos hacia la sombra del tejadillo de las cuadras.",
        "Allí estaba Yami Sukehiro, recostado contra una columna de madera con su cigarro humeante en la comisura de los labios, el torso ancho descubierto bajo una chaqueta abierta y su katana descansando en el hombro.",
        "—¡Y-Yami-san! —balbuceó Natsuki, poniéndose pálido al instante—. ¿D-Desde cuándo estás despierto?",
        "Yami ignoró la pregunta. Se acercó con pasos lentos y pesados, exhalando una densa bocanada de humo hacia el cielo matutino. Se detuvo justo frente a Natsuki, mirándolo de arriba abajo con sus ojos oscuros e imperturbables, hasta que apoyó una mano descomunal y pesada como plomo sobre el hombro del chico del chándal.",
        "Natsuki se encogió varios centímetros bajo el peso de esa mano.",
        "—Así que tienes oscuridad en las venas, novato —dijo Yami con una sonrisa torcida y peligrosa que dejaba ver sus dientes afilados—. La magia negra no es para señoritos que dibujan círculos mágicos en salones limpios. Es un elemento salvaje. Absorbe la luz, corta el espacio y no perdona a los tibios.",
        "—¿E-En serio...? —susurró Natsuki tragando saliva ruidosamente.",
        "—Si de verdad quieres dejar de ser alguien débil que solo sabe correr —continuó Yami, dándole un apretón en el hombro que hizo que a Natsuki se le salieran los ojos de las órbitas—, yo te enseñaré a empapar esa espada de madera con tus sombras.",
        "Una chispa de esperanza brilló en el rostro de Natsuki por una fracción de segundo.",
        "—¡¿De verdad, Yami?! ¡Muchas gracias, prometo ser—!",
        "—Pero te advierto una cosa —le cortó Yami, acercando su rostro endurecido por cicatrices al del muchacho—. Vas a sufrir como nunca en tu vida. Vas a entrenar hasta que las piernas se te duerman, vas a blandir esa espada hasta que los tendones te ardan, y cuando sientas que no puedes dar un paso más... vas a superar tus límites. Si no lo haces, te entierro vivo en el bosque.",
        "Un escalofrío recorrió la espalda de Natsuki de arriba abajo. Miró a Seier buscando auxilio, pero solo encontró mi asentimiento mudo y una leve sonrisa de aprobación.",
        "Natsuki apretó la empuñadura de madera con ambas manos, respiró hondo y, pese al terror evidente en sus ojos, sostuvo la mirada de Yami con una chispa de fuego inquebrantable.",
        "—...Acepto el trato, Yami. No pienso rendirme.",
        "Yami soltó una carcajada ronca, tiró la colilla al suelo y la pisó con su bota.",
        "—Bien dicho, idiota. Ahora ponte en guardia y lánzame cien estocadas antes de que termine mi próximo cigarro. ¡Muévete!",
        "La espada de Natsuki volvió a silbar en el aire con una energía renovada, mientras las primeras briznas de sombra comenzaban a titilar alrededor de su filo.",
        "En la Clase F-2, nadie se quedaba atrás."
    ]
}

def compilar_capitulo_13():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    from agregar_capitulo_12 import CAPITULO_12
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12, CAPITULO_13]
    
    # 1. Guardar texto plano del Capítulo 13
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "13_Capitulo_13.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_13['header'].upper()}: {CAPITULO_13['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_13['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 13
    docx_cap13 = os.path.join(NOVELA_DIR, "capitulos", "13_Capitulo_13.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_13['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_13['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_13['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap13)
    print(f"[OK] Word individual Capítulo 13 guardado en: {docx_cap13}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 13)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_13()
