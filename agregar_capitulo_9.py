# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_9 = {
    "header": "Capítulo 9",
    "title": "El Llamado del Bosque",
    "paragraphs": [
        "El resto de la segunda semana transcurrió con una calma productiva y bien aprovechada.",
        "Entre las clases de estrategia y las horas libres en los campos de práctica, me concentré en lo que más me hacía falta: afinar el flujo de mis dos afinidades elementales hacia mi cuerpo. Con los consejos prácticos de Ken y horas de repetición solitaria, mi progreso fue lento pero seguro.",
        "Ya no se trataba de una chispa aislada que me quemaba la ropa. Ahora era capaz de sostener un calor rojizo y constante en mis nudillos al momento de lanzar un golpe, y al mismo tiempo canalizar la pesadez de la tierra hacia mis piernas para arraigarme al suelo como un árbol milenario frente a los impactos.",
        "Todavía estaba lejos de conjurar proyectiles a distancia como los magos convencionales, pero para un peleador de corta distancia, ese blindaje interno marcaba una diferencia abismal.",
        "* * *",
        "El cambio drástico llegó el viernes por la tarde, cuando el profesor Vaneck mandó a cerrar las puertas del aula F-2 con gesto solemne.",
        "—Guarden sus cuadernos y presten atención —ordenó Vaneck, apoyando ambas palmas sobre el escritorio—. A partir del lunes de la tercera semana, comenzará su primera prueba oficial de campo.",
        "Un murmullo de anticipación recorrió los bancos.",
        "—Toda la clase pasará una semana completa internada en el Bosque de las Sombras Bajas —continuó el instructor—. Sin comodidades, sin sirvientes y sin raciones preparadas por la cocina de la academia. Tendrán que cazar, montar campamentos, montar guardias nocturnas y, lo más importante: explorar y limpiar un calabozo de nivel bajo ubicado en las profundidades del valle.",
        "Vaneck recorrió con sus ojos severos los rostros de cada uno de nosotros.",
        "—Muchos en las clases superiores creen que esto es solo una excursión para sumar puntos en la libreta. Se equivocan. Esta academia fue fundada para tiempos de crisis. Si el día de mañana la barrera demoníaca cede o estalla una guerra en las fronteras y el Rey emite el llamado a las armas, no pelearán junto a los aristócratas de la Clase A. Quienes estarán a su izquierda y a su derecha en las trincheras... son las personas sentadas en esta misma sala. Aprendan a confiar en ellos, aprendan sus nombres y cuiden sus espaldas. Porque de eso dependerá que vuelvan vivos.",
        "Las palabras de Vaneck calaron hondo. En un mundo donde mis propios padres habían muerto abandonados por la falta de refuerzos leales, el concepto de camaradería real tenía un peso sagrado.",
        "* * *",
        "Al sonar la campana, nuestro grupo habitual se reunió en las mesas del patio exterior para organizar los preparativos. Caesar, Lucy y Lighter no tardaron en acercarse para sumarse a la charla.",
        "—Una semana en el bosque... —murmuró Kaoruko, jugando nerviosa con la correa de su mochila—. Nunca he dormido al aire libre tantos días seguidos. Espero no ser una molestia para el equipo.",
        "—No digas tonterías, Kaoruko —le dijo Subaru con su tono firme y sereno—. Eres la sanadora principal de nuestra escuadra. En un calabozo, quien mantiene al equipo con vida eres tú. Yo me encargaré de armar un botiquín con ungüentos y vendas suficientes para todos.",
        "—¡Esa es la actitud! —exclamó Ken sonriendo con los brazos en la nuca—. ¡Además, las bestias de nivel bajo no serán rivales para nosotros! Si alguna se pasa de lista, le dejamos caer encima los puños de Seier, Lighter y los míos.",
        "—No te confíes, cabeza hueca —le advirtió Lucy cruzándose de brazos con aire autoritario—. Los calabozos no son arenas de sparring. Tienen trampas de presión, pasillos angostos donde no puedes pelear con comodidad y monstruos que emboscan en la oscuridad.",
        "—En eso Lucy tiene razón —asintió Lighter, acomodándose las gafas oscuras—. En los páramos de Calydon aprendimos por las malas que una bestia débil te puede arrancar una pierna si te descuidas montando guardia a las tres de la mañana.",
        "Caesar King apoyó una mano pesada y amistosa en mi hombro, mirándome a los ojos.",
        "—La F-2 tiene que dividirse en escuadras de reconocimiento y asalto, Seier. Tú y yo seremos la vanguardia pesada para frenar a cualquier criatura grande que intente romper la formación. ¿Podemos contar contigo?",
        "—Por supuesto —respondí con mi voz tranquila y segura—. Mi guardia no se va a quebrar.",
        "Kaoruko me miró desde el otro lado de la mesa, y una sonrisa dulce y aliviada iluminó su rostro.",
        "—Con ustedes al frente, siento que todo saldrá bien —dijo Kaoruko en voz baja.",
        "—Bueno, basta de charla —dijo Subaru poniéndose de pie con una libreta en mano—. Tenemos dos días para comprar cuerdas, pedernales, tiendas de campaña y raciones secas. Y más les vale a todos empacar ropa que no les importe manchar de barro.",
        "—¡A la orden, capitana Subaru! —bromeó Ken, arrancándole una mirada fulminante que nos hizo reír a todos.",
        "La primera gran prueba de la academia estaba a punto de comenzar. Y aunque el peligro en el bosque fuera real, la confianza forjada entre nosotros nos hacía sentir más fuertes que cualquier noble de la capital."
    ]
}

def compilar_capitulo_9():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9]
    
    # 1. Guardar texto plano del Capítulo 9
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "09_Capitulo_9.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_9['header'].upper()}: {CAPITULO_9['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_9['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 9
    docx_cap9 = os.path.join(NOVELA_DIR, "capitulos", "09_Capitulo_9.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_9['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_9['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_9['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap9)
    print(f"[OK] Word individual Capítulo 9 guardado en: {docx_cap9}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 9)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_9()
