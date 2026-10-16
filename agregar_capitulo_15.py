# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_15 = {
    "header": "Capítulo 15",
    "title": "La Pesadilla Teórica y las Aguas Termales",
    "paragraphs": [
        "El lunes comenzó con el retorno formal a las aulas y una escena digna de un hospital de campaña.",
        "Cuando entré al salón de la F-2, Subaru Natsuki estaba literalmente derretido sobre su pupitre. Tenía ojeras marcadas hasta los pómulos, los brazos colgando hacia el suelo como fideos cocidos y un hilo de saliva amenazando con escapar de su boca. Parecía un zombi recién desenterrado.",
        "—Sobreviví... apenas... pero sobreviví... —balbuceó Natsuki con voz ronca en cuanto me vio pasar.",
        "A su lado, Yami Sukehiro entró con las manos en los bolsillos, masticando su brizna de hierba habitual y con su katana al cinto.",
        "—Deja de quejarte, novato —dijo Yami, dándole una palmada en la espalda que hizo que a Natsuki le crujiera la columna entera—. Creí que te desmayarías al corte número cincuenta, pero aguantaste hasta el final. Tienes un potencial diminuto en esa cabeza hueca... pero al menos existe.",
        "Pese al dolor evidente en cada músculo de su cuerpo, los ojos de Natsuki brillaron por un instante con un orgullo infantil que no logró disimular.",
        "—¡Ja! ¿Lo escucharon? ¡El salvaje de la katana dijo que tengo potencial! —graznó Natsuki hacia el techo, antes de desplomarse de nuevo sobre la madera con un quejido agónico.",
        "* * *",
        "Diez minutos después, la puerta se abrió de par en par con un golpe seco.",
        "El profesor Vaneck entró con una pila enorme de carpetas amarillentas y una expresión que solo auguraba malas noticias. Dejó caer los papeles sobre el escritorio con un estruendo que hizo que toda la clase se enderezara por puro instinto.",
        "—Abran bien los oídos —anunció Vaneck sin preámbulos—. El período de gracia física ha terminado. En dos semanas comienzan los exámenes teóricos semestrales de la academia.",
        "Un silencio sepulcral se apoderó del aula.",
        "—Historia del reino, genealogía noble, geografía de maná y teoría avanzada de círculos de invocación elemental —enumeró el profesor con una frialdad cruel.",
        "El efecto fue instantáneo y devastador.",
        "La mitad de la clase —los que vivíamos para el combate cuerpo a cuerpo y la fuerza bruta— quedó petrificada como estatuas de sal. Ken abrió la boca con los ojos desorbitados, Renji se agarró la cabeza como si le fuera a estallar, Caesar King torció el gesto con una gota de sudor frío en la sien, y yo mismo sentí un nudo sordo en el estómago al recordar mi cero en la prueba de magia.",
        "En contraste, el bando intelectual reaccionó con total tranquilidad. Kaoruko ya tenía sus plumas alineadas y sus cuadernos impecables; Subaru Hoshina revisaba sus apuntes con una ceja alzada; Lucy sonreía con suficiencia cruzada de brazos; y Erza y Rukia asentían con la serenidad de quienes dominaban los manuales al dedillo.",
        "—Profesor... ¿tiene que ser tanta teoría? —se atrevió a balbucear Ken con voz temblorosa—. ¡Somos guerreros de la F, no monjes de biblioteca!",
        "—Un guerrero que no sabe leer un mapa ni calcular la resistencia mágica de una barrera es solo carne de cañón con espada —sentenció Vaneck con dureza—. Y más les vale estudiar como si sus vidas dependieran de ello, porque este año la dirección ha implementado un incentivo especial.",
        "El instructor se cruzó de brazos, barriendo la sala con su mirada penetrante.",
        "—Todas las clases cuyo promedio general quede de la mitad hacia arriba en la tabla de rendimiento de toda la academia —A-1, A-2, B-1, C-1, etcétera— recibirán una recompensa oficial: dos días completos con todos los gastos pagos en la Suite Real de Aguas Termales del balneario privado de la corona.",
        "La mención de las aguas termales cayó como un rayo en el aula.",
        "—¡¿Aguas termales reales?! —exclamó Burnice dando un brinco que casi tira su pupitre—. ¡Dicen que esas piscinas tienen maná curativo y que te sirven postres y bebidas de frutas flotando en bandejas de madera!",
        "—Comida ilimitada y saunas calientes para los músculos... —murmuró Lighter, bajándose las gafas oscuras con los ojos muy abiertos.",
        "—Un baño caliente después de meses de barro... —suspiró Kaoruko con las mejillas sonrojadas de anhelo.",
        "—No se hagan tantas ilusiones —advirtió Vaneck con una sonrisa irónica—. Históricamente, ese premio siempre se lo reparten las Clases A y B. Los hijos de los duques no tienen que romperse las manos aprendiendo a leer; ellos llevan tutores privados desde los cuatro años. Si la Clase F-2 quiere ese balneario... tendrán que hacer que hasta el último bruto de esta sala apruebe los exámenes.",
        "* * *",
        "En cuanto el profesor salió del aula, la campana sonó y el caos se desató.",
        "Ken, Renji y yo estábamos rodeados por un muro de miradas inquisidoras.",
        "—Muy bien, escuchen todos los descerebrados —dijo Lucy poniéndose de pie sobre su silla y apuntándonos con su gorra militar—. ¡No pienso perder esas aguas termales por culpa de un grupo de gorilas que no saben quién fue el tercer Rey de Solaria!",
        "—¡Apoyo la moción! —declaró Erza con aura marcial e implacable—. A partir de hoy, se cancelan los entrenamientos libres de la tarde. Formaremos grupos de estudio obligatorios. Quien intente escaparse para ir a pelear... sentirá el filo de mi espada de madera.",
        "—Piedad... por favor, piedad... —gimió Natsuki desde su pupitre.",
        "Kaoruko se acercó a mi mesa con una pila de resúmenes prolijos y una sonrisa tierna pero decidida.",
        "—No te preocupes, Seier-kun —me dijo en voz baja, inclinándose hacia mí con aroma a flores—. Yo te ayudaré con la teoría de historia y los círculos mágicos. Si repasamos todos los días, ¡seguro que lo logramos juntos!",
        "La miré a los ojos, notando el brillo dulce y entusiasta en su mirada.",
        "—Trato hecho, Kaoruko —respondí, esbozando una media sonrisa resigneda—. Prefiero pelear contra tres gólems en el bosque... pero por esas aguas termales y tus apuntes, me aprendo el libro entero.",
        "La guerra de los libros en la Clase F-2 acababa de comenzar."
    ]
}

def compilar_capitulo_15():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    from agregar_capitulo_12 import CAPITULO_12
    from agregar_capitulo_13 import CAPITULO_13
    from agregar_capitulo_14 import CAPITULO_14
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12, CAPITULO_13, CAPITULO_14, CAPITULO_15]
    
    # 1. Guardar texto plano del Capítulo 15
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "15_Capitulo_15.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_15['header'].upper()}: {CAPITULO_15['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_15['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 15
    docx_cap15 = os.path.join(NOVELA_DIR, "capitulos", "15_Capitulo_15.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_15['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_15['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_15['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap15)
    print(f"[OK] Word individual Capítulo 15 guardado en: {docx_cap15}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 15)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_15()
