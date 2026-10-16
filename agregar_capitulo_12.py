# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_12 = {
    "header": "Capítulo 12",
    "title": "Notas Imperfectas y Lazos Forjados",
    "paragraphs": [
        "El lunes por la mañana, la atmósfera en el aula F-2 era de una tensa expectación.",
        "El profesor Vaneck entró con su paso marcial habitual, dejó caer una carpeta de cuero sobre el pupitre y barrió el salón con su ojo crítico. Todos los asientos estaban llenos; nadie faltaba y nadie lucía las secuelas de una derrota.",
        "—Tengo en mis manos las actas de evaluación de la expedición de supervivencia —comenzó Vaneck, cruzando las manos en la espalda—. Para sorpresa de los archimaestres del consejo, la Clase F-2 no solo completó los siete días sin solicitar una sola bengala de rescate, sino que limpió el Calabozo de Piedra Gris hasta su cámara final.",
        "Ken estuvo a punto de soltar un grito de victoria, pero una mirada gélida del profesor lo sentó de golpe en la madera.",
        "—Sin embargo —continuó Vaneck enarcando una ceja—, su calificación final es de un ocho sobre diez. Muy superior a la media, pero lejos de la perfección.",
        "—¿Por qué, profesor? —preguntó Erza con la espalda recta como una tabla, claramente mortificada por no haber alcanzado la nota máxima—. Cumplimos todos los objetivos tácticos y no sufrimos bajas.",
        "—Porque la academia también califica ordenanza y preservación —respondió Vaneck, abriendo el informe—. Primero: una deducción del diez por ciento por daño estructural severo en el tercer nivel. Quemar las paredes con combustible inflamable para calentar un gólem no figura en ningún manual de la iglesia.",
        "Burnice soltó una risita culpable mientras Lucy le propinaba un codazo en las costillas.",
        "—Segundo —prosiguió el instructor mirando hacia el rincón—: la bitácora de campo entregada por Akame y Natsuki tenía manchas sospechosas de grasa de jabalí en los márgenes.",
        "Akame desvió la vista hacia la ventana con total inocencia, mientras Subaru Natsuki se encogía de hombros con una sonrisa nerviosa.",
        "—Y tercero: ni uno solo de los monstruos fue abatido mediante círculos de magia canónicos de tercer orden. Pura violencia física, golpes de tierra, cortes de vacío y fuego marcial. Los examinadores de la corte dijeron que parecían una horda de mercenarios más que cadetes de Solaria.",
        "Vaneck cerró la carpeta con un golpe seco, y por primera vez en semanas, una sonrisa tenue y orgullosa asomó en su rostro severo.",
        "—Pero volvieron enteros. Y para mí, eso vale más que cualquier diez en un pergamino dorado. La dirección les ha concedido dos días libres de recuperación. Disfrútenlos.",
        "En cuanto Vaneck salió del aula, el salón estalló en vítores y palmadas.",
        "* * *",
        "Los dos días de descanso sirvieron para que la tensión de las armas diera paso a una camaradería genuina y cotidiana.",
        "El martes por la tarde, Kaoruko me esperó a la salida de los dormitorios y me invitó a subir a la terraza del viejo invernadero botánico, un rincón apartado del campus rodeado de helechos gigantes y flores de maná azuladas.",
        "—Preparé esto esta mañana en la cocina común —dijo Kaoruko con las mejillas encendidas de un carmesí suave, abriendo una cesta de tela—. Son bollitos de manzana con miel y canela... quería agradecerte por lo del calabozo. Cuando ese gólem levantó el puño hacia mí... de verdad creí que no lo contaba.",
        "Tomé uno de los bollos dorados. Aún estaba tibio al tacto. Le di un bocado generoso y el sabor dulce, especiado y suave me hizo cerrar los ojos de gusto.",
        "—Están mejores que los de la pastelería del centro —le dije con total sinceridad, mirándola a los ojos.",
        "A Kaoruko casi se le cae la servilleta de las manos por la emoción. Juntó las palmas frente a su pecho, sonriendo con tanta felicidad que sus enormes ojos azules parecían dos zafiros iluminados por el sol.",
        "—¡Waaah, qué alivio! —suspiró ella con una risita tímida—. Temía que la masa hubiera quedado muy seca. Seier-kun... eres muy amable conmigo. A veces la gente te juzga por lo callado o grande que eres, pero yo sé que tienes un corazón enorme.",
        "—Solo cuido a los míos, Kaoruko —respondí con voz baja y tranquila, sentándome en el pretil de piedra a su lado—. En la frontera me enseñaron que la fuerza solo sirve si protege lo que de verdad importa.",
        "Kaoruko se sentó junto a mí, tan cerca que sentía el calor de su hombro y el suave aroma a flores que desprendía su cabello negro. Pasamos la tarde entera conversando entre risas y bocados, dejando que una cercanía dulce y reconfortante floreciera entre nosotros sin necesidad de prisas.",
        "* * *",
        "En los patios inferiores, los demás también encontraban su propio ritmo.",
        "Junto a los lavaderos, los dos Subaru estaban enzarzados en una discusión absurda.",
        "—¡Te digo que yo debería ser 'Subaru Mayor' porque soy un año mayor que tú! —protestaba Natsuki agitando los brazos.",
        "—Y yo te digo que te llamaré simplemente Natsuki si no quieres que te congele los calcetines —respondía Subaru Hoshina con su expresión fría y distante, aunque le entregó un frasco de ungüento para las ampollas de las manos—. Tómalo. Si no te curas los dedos, no podrás cargar las mochilas en la próxima práctica.",
        "—Vaya, mira quién tiene un corazoncito bajo esa armadura de hielo —bromeó Natsuki, ganándose una mirada fulminante que lo hizo reír a carcajadas.",
        "Más allá, en el patio trasero de los establos, se había formado lo que Ken bautizó como 'El Club de la Carne'.",
        "Yami, Renji y Akame habían armado una parrilla improvisada con piedras refractarias, asando las costillas de jabalí que habían traído del bosque. Akame cortaba la carne con una velocidad y precisión quirúrgicas, comiendo trozo tras trozo con una expresión de beatitud absoluta, mientras Yami fumaba recostado contra una carreta y Renji intentaba avivar las brasas a base de bufidos de maná.",
        "—¡Oye, Renji, no le metas tanto viento o vas a quemar la grasa, idiota! —le gritaba Yami sin abrir los ojos.",
        "—¡Cállate, pesado, que este fuego lo controlo yo! —respondía Renji con una vena marcada en la frente.",
        "Cerca de allí, Erza Scarlet, Caesar King y Lighter compartían anécdotas de armas sentados en las gradas. Lighter le mostraba a Ken cómo pivoter sobre el talón para esquivar estocadas rectas, mientras Caesar y Erza debatían con seriedad sobre el centro de gravedad de una espada de dos manos frente a un hacha de guerra.",
        "Miré a mi alrededor mientras el sol teñía de naranja los tejados de la academia.",
        "Ninguno de nosotros venía de palacios de cristal. Todos cargábamos cicatrices, pérdidas o el estigma de ser considerados descartes por la corte. Pero al vernos reír, entrenar y compartir la comida, comprendí que la nobleza de sangre no significaba nada.",
        "La verdadera nobleza se forjaba en el barro, en la lealtad y en la sangre compartida.",
        "Y la Clase F-2 estaba lista para cualquier tormenta que se avecinara."
    ]
}

def compilar_capitulo_12():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12]
    
    # 1. Guardar texto plano del Capítulo 12
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "12_Capitulo_12.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_12['header'].upper()}: {CAPITULO_12['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_12['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 12
    docx_cap12 = os.path.join(NOVELA_DIR, "capitulos", "12_Capitulo_12.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_12['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_12['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_12['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap12)
    print(f"[OK] Word individual Capítulo 12 guardado en: {docx_cap12}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 12)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_12()
