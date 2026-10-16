# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_3 = {
    "header": "Capítulo 3",
    "title": "La Clase F-2 y la Chica de los Dulces",
    "illustration": os.path.join(NOVELA_DIR, "imagenes", "kaoruko_waguri.png"),
    "paragraphs": [
        "El sistema de la Gran Academia Militar y Mágica de Solaria estaba estructurado como una pirámide despiadada.",
        "Los estudiantes no eran divididos únicamente por su año de ingreso, sino por un escalafón estricto de seis rangos: desde la élite absoluta de la Clase A hasta el peldaño inferior de la Clase F. A su vez, cada letra se subdividía en aulas numéricas según el promedio exacto de admisión: A-1, A-2, B-1... hasta F-1 y F-2.",
        "Los herederos de las grandes casas ducales, los prodigios con afinidad de maná pura y los favoritos del clero ocupaban las aulas doradas de la Clase A, situadas en las torres más altas del castillo central, con alfombras de lana, báculos de cristal y balcones con vista a los jardines imperiales. Allí era donde residía gente como Elysea de Vane.",
        "A mí, por mi parte, me habían asignado a la Clase F-2.",
        "Mi puntuación perfecta en la prueba física y la victoria fulminante en el duelo habían sido neutralizadas por el cero rotundo en conjuración de hechizos. Sumado a que la Casa Silva era considerada un linaje arruinado y sin respaldo político en la corte, la administración no dudó un segundo en enviarme al fondo del barril.",
        "El ala F estaba ubicada en la planta baja del antiguo bastión oeste. Los pasillos eran de piedra gris sin pulir, iluminados por antorchas de hierro en lugar de lámparas de cristal encantado, y las ventanas daban a los patios de armas secundarios.",
        "Sin embargo, al recorrer aquellos corredores, el ambiente me resultó infinitamente más agradable que el aire enrarecido y falso de la plaza central.",
        "Aquí no había nobles arrogantes mirándose al espejo ni damas con risitas fingidas. La Clase F-2 era el hogar de los marginados: vástagos de familias nobles caídas en desgracia, soldados que habían ascendido por mérito propio y plebeyos becados que llevaban los codos gastados de estudiar y las manos curtidas de trabajar. Gente que no tenía el lujo de heredar un título y que estaba dispuesta a partirse el alma para ser alguien.",
        "Ese era mi terreno.",
        "* * *",
        "Iba caminando hacia la puerta de madera maciza con la placa de bronce grabada con 'F-2', con mi mochila al hombro y el paso pesado de mis noventa kilos de músculo, cuando al doblar la esquina de un arco de piedra sentí un impacto repentino contra mi costado.",
        "Para mí, con mi centro de gravedad bajo y la estabilidad arraigada de un luchador de Muay Thai, fue apenas un roce que ni siquiera movió mis hombros. Pero quien venía en sentido contrario no tuvo la misma suerte.",
        "—¡A-Ah! —un grito ahogado y agudo resonó en el pasillo.",
        "Una pila de gruesos tomos encuadernados en cuero y rollos de pergamino voló por los aires. La persona que los cargaba perdió el equilibrio y cayó sentada sobre las losas de piedra con un ruido sordo.",
        "Me detuve en seco y me giré de inmediato.",
        "En el suelo yacía una chica menuda, que parecía haber chocado contra un muro de granito. Se frotaba la cadera con una mano enguantada mientras miraba con pánico sus notas desparramadas por todo el suelo.",
        "Bajo mi aspecto tosco y mi costumbre de hablar con pocas palabras, jamás me ha gustado ser un desconsiderado.",
        "—Disculpa —dije con mi voz profunda y rasposa, agachándome de inmediato frente a ella—. Iba distraído. ¿Te hiciste daño?",
        "La chica alzó la vista con timidez, y por un instante el pasillo pareció quedarse en silencio.",
        "Era menuda y de aspecto frágil, con una belleza dulce y conmovedora que contrastaba salvajemente con la frialdad aristocrática de Elysea. Tenía una larga y sedosa melena negra que caía en ondas suaves sobre sus hombros, enmarcando un rostro angelical de piel nívea. Pero lo más cautivador eran sus ojos: dos orbes grandes y luminosos de un azul zafiro brillante, cargados de una expresividad inocente que parecía no tener malicia alguna.",
        "En sus mejillas ardía un rubor encendido, parte por la vergüenza de la caída y parte por encontrarse de repente frente a un tipo con hombros como dos rocas y cara de pocos amigos.",
        "—N-No... no se preocupe... fue culpa mía por ir corriendo sin mirar —respondió con una vocecita suave y temblorosa, como el tintineo de una campanilla de cristal.",
        "Extendí mi mano derecha hacia ella. Era una mano ancha, pesada, con los nudillos endurecidos por años de golpear madera y hierro.",
        "—Toma. Apóyate —dije con calma.",
        "La chica dudó una fracción de segundo al ver el tamaño de mi mano, pero al percibir que no había la menor hostilidad en mi mirada, colocó su mano pequeña y tibia sobre mi palma. Sus dedos eran delgados y suaves como pétalos de rosa. Con un movimiento firme y controlado, sin brusquedad, tiré de ella hacia arriba y la ayudé a ponerse en pie.",
        "Luego me arrodillé sobre un solo pie para recoger sus libros y pergaminos desparramados.",
        "Ella se apresuró a agacharse a mi lado, juntando hojas a toda velocidad con las mejillas todavía teñidas de carmesí. Noté que sus dedos temblaban levemente por los nervios.",
        "—De verdad... muchas gracias... —murmuró, abrazando los tomos recuperados contra su pecho.",
        "La miré de reojo mientras me ponía de pie. Su uniforme básico de la academia le quedaba ligeramente holgado en los hombros, acentuando su aire delicado.",
        "—Siento haberte tirado —repetí, rascándome la nuca con cierta tosquedad—. Tengo la costumbre de caminar como un buey. Para compensarte... después de la primera clase, ¿te gustaría algo dulce? Compré una caja de pastelillos de miel y frutos rojos en el mercado central antes de entrar. Son bastante buenos.",
        "En cuanto la palabra 'dulce' salió de mis labios, ocurrió una transformación asombrosa.",
        "La timidez y el miedo que envolvían a la chica desaparecieron en un parpadeo. Sus enormes ojos azules se abrieron como platos, encendiéndose con un brillo centelleante de pura fascinación y deleite, como si le hubiera ofrecido el mayor tesoro del imperio.",
        "—¿D-Dulces...? ¿Pastelillos de frutos rojos? —preguntó, tragando saliva de manera adorable y asintiendo con la cabeza a una velocidad cómica—. ¡S-Sí, por favor! Digo... si no es una molestia... ¡me encantaría!",
        "La espontaneidad de su reacción fue tan genuina que una media sonrisa involuntaria se dibujó en mis labios.",
        "—No es molestia. Me llamo Seier. Seier Silva.",
        "Ella juntó las manos frente a su pecho, sonriendo con una timidez cálida y radiante que iluminó todo el pasillo.",
        "—Yo soy Kaoruko Waguri. Un placer conocerte, Seier-kun... ¿también vas a la Clase F-2?",
        "—Sí. Parece que seremos compañeros de aula.",
        "Los ojos de Kaoruko brillaron aún más, aliviada de no estar sola en su primer día.",
        "* * *",
        "Entramos juntos al aula F-2 y nos sentamos en una de las bancas de roble de la segunda fila.",
        "El salón era espacioso y austero. Los cerca de treinta alumnos que llenaban los pupitres eran variados: desde un muchacho con brazos de herrero hasta jóvenes nobles empobrecidos que repasaban nerviosos sus apuntes.",
        "Poco después, la puerta se abrió con un chasquido autoritario.",
        "Entró el profesor de la materia: el Maestro Vaneck, un erudito de cabello entrecano, túnica marrón sencilla y una cicatriz profunda en el antebrazo izquierdo que delataba su pasado como capellán militar. Dejó una pesada pila de atlas históricos sobre la mesa y barrió la sala con ojos penetrantes.",
        "—Bienvenidos a la Clase F-2 —dijo Vaneck con voz seca y resonante—. Si esperan que este lugar sea un asilo de mediocres donde relajarse porque los nobles de la Clase A los consideran basura, pueden marcharse por donde entraron. Aquí se forjan los que no tienen derecho a equivocarse. Mi nombre es Vaneck, y hoy comenzaremos con Teoría e Historia del Reino de Solaria.",
        "Vaneck desplegó un gran mapa cartográfico sobre la pizarra de pizarra negra. El mapa mostraba la geografía del continente: el núcleo protegido de la Capital Imperial, los grandes valles fértiles de las familias ducales y, rodeando toda la frontera meridional, una franja oscura y montañosa marcada con runas rojas: *El Desfiladero de las Sombras*.",
        "—Hace doscientos años —comenzó Vaneck, paseándose entre las filas—, la barrera primordial que separa nuestro plano del Abismo Demoníaco sufrió su primera gran fractura. Los libros de historia que leen los señoritos de la Clase A dicen que el Rey y los paladines de la iglesia salvaron el mundo con oraciones y espadas sagradas.",
        "Vaneck golpeó la madera del mapa con los nudillos.",
        "—Una mentira piadosa para los salones de té. La realidad es que el reino sigue en pie porque hubo familias de vanguardia en la frontera sur que pusieron sus pechos de hierro contra las garras de las bestias. Hombres y mujeres que combatieron en el fango, que no recibieron refuerzos a tiempo y que se sacrificaron para que las familias de la capital pudieran dormir tranquilas.",
        "Al escuchar esas palabras, un nudo sordo se apretó en mi garganta.",
        "No pude evitar apretar el medallón de bronce con el león de los Silva que colgaba bajo mi camisa de lino.",
        "A mi lado, Kaoruko estaba tomando notas con una caligrafía primorosa y veloz. Pero al escuchar la mención de las familias fronterizas, detuvo su pluma por un instante. Giró su rostro hacia mí, mirándome de reojo con esos ojos zafiro colmados de una empatía y un respeto silenciosos. Claramente había reconocido el apellido 'Silva' en cuanto me presenté.",
        "Nuestras miradas se encontraron durante unos segundos. No hizo preguntas indiscretas ni mostró lástima barata; simplemente me dedicó una sonrisa suave y comprensiva, transmitiéndome una calidez que me reconfortó más de lo que jamás admitiría en voz alta.",
        "* * *",
        "La clase de dos horas pasó volando entre tácticas antiguas de asedio, geografía de maná y las leyes feudales del reino.",
        "Cuando la campana de bronce del patio repicó anunciando el final de la sesión matutina, el profesor Vaneck recogió sus pergaminos y salió del aula. El bullicio de los alumnos comenzó a llenar el salón mientras guardaban sus plumas.",
        "Me giré hacia Kaoruko. La chica ya tenía sus libros perfectamente apilados y me miraba con una expresión de anticipación tan tierna que parecía a punto de dar saltitos en su asiento.",
        "—Lo prometido es deuda —dije, bajando mi mochila y abriendo la solapa de cuero.",
        "Saqué una cajita de madera de abedul pulida y la coloqué sobre el pupitre entre los dos. Al abrir la tapa, un aroma embriagador a mantequilla fresca, miel tostada y frutos del bosque inundó el aire a nuestro alrededor.",
        "Dentro había cuatro pastelillos esponjosos, cubiertos con una capa brillante de glaseado carmesí y decorados con crema batida blanca.",
        "—¡Waaaah...! —un suspiro de genuino éxtasis escapó de los labios de Kaoruko. Sus mejillas se tiñeron de un rosa intenso y sus pupilas brillaron como dos estrellas—. ¡S-Se ven deliciosos! ¿De verdad puedo tomar uno?",
        "—Toma los que quieras. Para eso son —respondí.",
        "Kaoruko juntó sus manos en una breve plegaria de agradecimiento, tomó uno de los pastelillos con la punta de sus dedos delicados y le dio un bocado pequeño y cuidadoso.",
        "En cuanto la crema y la masa tocaron su paladar, cerró los ojos y soltó un quejido suave y placentero que hizo que el corazón me diera un vuelco involuntario en el pecho.",
        "—¡Mmmh...! ¡Está... increíblemente rico! —exclamó, llevándose una mano a la mejilla mientras masticaba con una sonrisa de felicidad absoluta—. La masa es tan suave... y los frutos tienen el punto exacto de acidez. ¡Nunca había probado algo tan bueno!",
        "La vi comer con tanto entusiasmo y devoción que no pude evitar que una risa baja y tranquila brotara de mi garganta.",
        "—Me alegra que te gusten, Waguri. Los compré en un callejón detrás de la plaza.",
        "—¡Puedes llamarme Kaoruko, Seier-kun! —dijo con una sonrisa radiante, limpiándose una diminuta mancha de glaseado de la comisura de los labios con la lengua, un gesto completamente inocente pero que desprendía un magnetismo dulce y peligroso a la vez—. De verdad... muchas gracias. Estaba tan asustada esta mañana... creí que todos en la academia serían fríos o crueles.",
        "—No todos —le dije, mirándola con firmeza—. Y en esta clase, mientras yo esté cerca, nadie va a meterse contigo.",
        "Kaoruko me miró fijamente. El sonrojo en sus mejillas se profundizó, pero esta vez no apartó la vista. En su mirada había algo más que gratitud: un destello de confianza plena y una atracción naciente que comenzaba a florecer entre los dos.",
        "En las alturas de la academia, los nobles conspiraban y las espadas de hielo afilaban su orgullo. Pero aquí abajo, en el aula olvidada de la F-2, había encontrado algo por lo que valía la pena pelear."
    ]
}

def compilar_capitulo_3():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3]
    
    # 1. Guardar texto plano
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "03_Capitulo_3.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_3['header'].upper()}: {CAPITULO_3['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_3['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word exclusivo Capítulo 3 (con ilustración si existe)
    docx_cap3 = os.path.join(NOVELA_DIR, "capitulos", "03_Capitulo_3.docx")
    
    # Creamos documento individual con la imagen de Kaoruko
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # Título
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_3['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_3['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)
    
    # Ilustración
    if os.path.exists(CAPITULO_3['illustration']):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(15)
        p_img.paragraph_format.space_after = Pt(4)
        r_img = p_img.add_run()
        r_img.add_picture(CAPITULO_3['illustration'], width=Inches(3.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(20)
        r_cap = p_cap.add_run("Ilustración: Kaoruko Waguri (Clase F-2)")
        r_cap.font.name = "Georgia"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(120, 100, 140)

    # Párrafos
    for p_text in CAPITULO_3['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- "):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap3)
    print(f"[OK] Word individual Capítulo 3 guardado en: {docx_cap3}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 + Cap 2 + Cap 3)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_3()
