# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_6 = {
    "header": "Capítulo 6",
    "title": "Chispas de Fuego y una Invitación",
    "paragraphs": [
        "El sábado por la mañana, los terrenos de la academia amanecieron en una calma casi total.",
        "La mayoría de los estudiantes aprovechaba el fin de semana para dormir hasta tarde o pasear por el distrito comercial de la capital, pero quedarse en la cama no iba conmigo. Tras la derrota por la mínima contra Ken el viernes, la necesidad de dominar el flujo elemental en mis puños me quemaba por dentro.",
        "Me calcé las botas, me puse unos pantalones cómodos de tela resistente y una camiseta sin mangas, y me dirigí hacia los campos de entrenamiento secundarios en los límites del bosque este.",
        "* * *",
        "Cuando llegué a la zona de claros rocosos, el eco de golpes secos y el chasquido inconfundible de llamas rompió el silencio de los árboles.",
        "A unos cincuenta metros, Ken estaba saltando entre tres troncos de entrenamiento, soltando combinaciones de puñetazos ardientes al aire con una energía envidiable.",
        "—¡Vaya, mira a quién tenemos aquí! —gritó Ken en cuanto me vio, deteniéndose en seco con una enorme sonrisa y limpiándose el sudor con el brazo—. ¡No me digas que las piernas ya no te duelen de ayer, Seier!",
        "—Duermo bien y me recupero rápido —respondí, caminando hacia él—. Veo que tú tampoco sabes quedarte quieto un sábado.",
        "—¡Para nada! Un artista marcial que descansa cuando sus rivales entrenan está destinado a perder —dijo Ken, lanzando un puño al aire—. ¿Viniste a trabajar en tu maná elemental?",
        "—Sí. Ayer quedó claro que pelear solo con el cuerpo tiene un límite. Si no logro despertar el fuego o la tierra en mis golpes, me voy a quedar estancado.",
        "Ken asintió con seriedad, cruzándose de brazos.",
        "—Mi maestra Hwa Young siempre me decía que los magos de academia se equivocan al pensar en la magia como fórmulas y círculos dibujados en el aire. Para los luchadores, el maná no es un cántico... es una extensión de la respiración.",
        "—¿Cómo lo haces tú?",
        "—Mira —Ken extendió su palma derecha hacia adelante—. Tú tienes fuego y tierra, ¿verdad? El fuego es explosión y emoción; la tierra es peso y arraigo. Para el fuego, no intentes sacarlo de tu piel como si lanzaras una flecha. Siente el calor en tu estómago, hazlo subir por el pecho hasta los hombros, y solo cuando el puño esté a un centímetro de impactar... déjalo detonar.",
        "Me coloqué en mi postura básica frente a una roca de tamaño mediano.",
        "Cerré los ojos y respiré hondo. En lugar de empujar el maná hacia afuera como un tonto, seguí la indicación de Ken: concentré la energía en el núcleo, la guié a través del hombro y solté un directo seco de derecha.",
        "¡FWOOSH!",
        "Una llamarada anaranjada, pequeña pero viva, brotó de mis nudillos y envolvió la superficie de la roca antes de apagarse en una estela de humo negro.",
        "Ken abrió los ojos como platos y soltó una carcajada ruidosa.",
        "—¡Eso es! ¡Lo lograste al primer intento, bestia! ¿Ves que no era tan difícil?",
        "—Apenas fue una chispa —dije, mirando mi puño humeante con una sonrisa satisfecha—. Pero ya entiendo la sensación.",
        "—¡Pues a repetir hasta que te salga natural! ¡Vamos a hacer sparring técnico!",
        "* * *",
        "Pasamos las dos horas siguientes intercambiando golpes controlados, corrigiendo la postura y puliendo la transición entre el peso de la tierra para bloquear y el calor del fuego para contraatacar. Estábamos empapados de sudor, con moretones y raspones cubriéndonos los brazos, cuando unas voces conocidas nos interrumpieron desde el sendero.",
        "—¡Lo sabía! ¡Estaban aquí! —la voz alegre de Kaoruko resonó entre los árboles.",
        "Ken y yo bajamos la guardia y nos giramos.",
        "Kaoruko venía caminando apresurada con una canasta de mimbre en una mano y una sonrisa radiante. A su lado, Subaru caminaba con su paso tranquilo y distante, sosteniendo una toalla limpia sobre el hombro y mirándonos con un leve suspiro.",
        "—¿Ustedes dos no saben hacer otra cosa que no sea golpearse? —preguntó Subaru con su habitual ceño fruncido, aunque dejó una jarra de té frío sobre un tocón de madera.",
        "—¡Hola, chicas! —saludó Ken con su entusiasmo habitual—. ¡Llegan justo a tiempo! ¡Seier acaba de encender fuego en sus puños!",
        "—¿En serio? ¡Eso es genial, Seier-kun! —exclamó Kaoruko, con los ojos brillando de admiración. Luego se fijó en nuestros brazos y su expresión cambió a una de maternal regaño—. Pero miren cómo quedaron... llenos de quemaduras leves y rasguños. Siéntense los dos ahora mismo.",
        "Ken y yo nos sentamos obedientemente en la hierba.",
        "Kaoruko se arrodilló frente a mí y colocó sus manos tibias sobre mi antebrazo derecho. Una luz azulada y fresca brotó de sus palmas, envolviendo la piel irritada. La sensación de alivio fue inmediata, como sumergir el brazo en un manantial helado.",
        "—Tienes una magia muy suave, Kaoruko —le dije en voz baja mientras la observaba trabajar tan concentrada.",
        "Ella alzó la vista un instante y sus mejillas se sonrojaron de inmediato.",
        "—M-Me alegra que te ayude... mientras pueda ser útil para ustedes, haré lo que sea.",
        "A nuestro lado, Subaru ayudaba a Ken a vendarse una muñeca, usando un hechizo menor de viento para refrescar el ambiente.",
        "—No te acostumbres, cabeza azul —le dijo Subaru con frialdad fingida a Ken—. No soy enfermera particular.",
        "—¡Jajaja, gracias de todos modos, Subaru! ¡Eres mucho más blanda de lo que aparentas!",
        "—Cállate si no quieres que te apriete más el vendaje.",
        "El ambiente entre los cuatro era tan natural y relajado que el cansancio del entrenamiento pareció desvanecerse por completo.",
        "* * *",
        "Cuando terminaron de curarnos, me puse de pie y me sacudí la tierra del pantalón.",
        "—Ken me ayudó a entender la canalización del fuego, y ustedes dos vinieron hasta aquí a curarnos y traernos agua fresca —dije, mirando a los tres con mi calma habitual—. Así que no voy a dejarlos ir con el estómago vacío. Yo invito el almuerzo.",
        "Ken se levantó de un salto como si le hubieran inyectado adrenalina pura.",
        "—¡¿En serio?! ¡Carne asada! ¡Dime que vamos por carne asada!",
        "—¡Pastelillos de la plaza! —añadió Kaoruko con las manos juntas y los ojos centelleantes de emoción infantil.",
        "Subaru intentó mantener su compostura seria, pero al ver las caras de los otros dos, una sonrisa tenue y casi invisible asomó en la comisura de sus labios.",
        "—Supongo que no puedo negarme si todos van —murmuró Subaru.",
        "—Está decidido —dije, echándome la mochila al hombro—. Vamos a la ciudad. Invito yo.",
        "Y entre risas, quejas de Ken sobre el hambre y la emoción de Kaoruko por los postres, los cuatro tomamos el camino hacia la puerta principal de Solaria, consolidando un grupo que cada día se sentía más unido."
    ]
}

def compilar_capitulo_6():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6]
    
    # 1. Guardar texto plano del Capítulo 6
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "06_Capitulo_6.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_6['header'].upper()}: {CAPITULO_6['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_6['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 6
    docx_cap6 = os.path.join(NOVELA_DIR, "capitulos", "06_Capitulo_6.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_6['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_6['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_6['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap6)
    print(f"[OK] Word individual Capítulo 6 guardado en: {docx_cap6}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 + Cap 2 + Cap 3 + Cap 4 + Cap 5 + Cap 6)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_6()
