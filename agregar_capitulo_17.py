# -*- coding: utf-8 -*-
import os
import sys

NOVELA_DIR = r"C:\Users\fsilv\OneDrive\Documentos\GitHub\Novela"
sys.path.insert(0, NOVELA_DIR)

from generar_libro import create_novel_document
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CAPITULO_17 = {
    "header": "Capítulo 17",
    "title": "El Milagro de las Aguas Termales",
    "paragraphs": [
        "El Muro de Cristal de la plaza mayor proyectaba las listas generales bajo un resplandor dorado que congregaba a cientos de estudiantes nerviosos.",
        "La tabla clasificaba a las veinticuatro aulas de la academia según su nota media global. La regla era implacable: solo las primeras doce clases —la mitad superior exacta— tendrían acceso a la legendaria Suite Real de Aguas Termales del balneario de la corona.",
        "—¡Abran paso, dejen ver! —gritaba Renji abriéndose camino a codazos entre la multitud aristocrática junto a Ken.",
        "Natsuki se puso de puntillas, temblando como una hoja, mientras Lucy y Erza revisaban fila por fila con ojos de águila.",
        "—Clase A-1: primer lugar... A-2: segundo... B-1: tercero... —murmuraba Lucy leyendo a toda prisa—. Noveno: C-2... Décimo: D-1... Undécimo: D-3...",
        "El corazón de toda la F-2 se detuvo.",
        "—Puesto doce... —Lucy se frotó los ojos y se pegó al cristal con la boca abierta—. ¡PUESTO DOCE: CLASE F-2!",
        "Silencio.",
        "Y luego, un rugido ensordecedor que hizo que los pájaros de las torres salieran volando.",
        "—¡CLASIFICAMOS! ¡ENTRAMOS DE PURA SUERTE EN EL ÚLTIMO LUGAR! —bramó Ken alzando a Natsuki por los aires como si fuera un muñeco de trapo.",
        "—¡Por dos centésimas de punto! ¡Superamos a la E-1 por un suspiro! —chilló Burnice abrazando a Caesar King entre saltos histéricos.",
        "A pocos metros, varios nobles de las clases intermedias nos miraban pálidos y boquiabiertos, sin poder creer que un aula repleta de plebeyos y marginados les hubiera arrebatado el billete a las aguas termales por los pelos.",
        "Kaoruko se giró hacia mí con los ojos llorosos de pura alegría, y sin pensarlo dos veces se abalanzó a abrazarme por el cuello. Su cuerpo menudo y suave se apretó contra mi pecho; olía a jazmines y canela. Le devolví el abrazo con una risa baja, sintiendo el calor de sus mejillas pegadas a mi hombro.",
        "—¡Lo logramos, Seier-kun! ¡De verdad lo logramos! —susurró con una emoción tan pura que me sacudió por dentro.",
        "* * *",
        "El viaje al balneario imperial en tres carruajes contratados por la academia fue una fiesta sobre ruedas.",
        "Para cuando cruzamos los portones de madera de cedro de la Suite Real, el cansancio acumulado de semanas de estudio y barro se desvaneció al instante. El lugar era un paraíso tallado en las faldas de una montaña volcánica: pabellones de madera aromática, estanques de agua termal humeante rodeados de piedras blancas lisas, linternas de papel de arroz y jardines de bambú.",
        "—Esto supera cualquier cama de Calydon —dijo Piper Wheel, descalzándose de inmediato y dejándose caer como un trapo sobre los tapices de plumas de la terraza—. No me despierten hasta el año que viene.",
        "Pulchra, la chica zorro de los Hijos de Calydon que casi siempre se mantenía en guardia y con la máscara puesta, se desató la tela por primera vez frente a nosotros. Reveló un rostro felino hermoso, de rasgos finos y ojos ámbar traviesos. Sus orejas peludas y su cola esponjosa se mecían de lado a lado con deleite mientras olfateaba el vapor del estanque.",
        "—Agua caliente con minerales de maná... —murmuró Pulchra con una sonrisa relajada, estirando sus brazos ágiles—. Es perfecta para quitarse el polvo de los caminos. Y miren allá... hay brochetas de salmón a la parrilla en el comedor.",
        "Akame apareció a su lado como un fantasma en cuanto escuchó la palabra 'salmón'.",
        "—Salmón de río fresco. Carne asada en piedra volcánica. Es un deber gastronómico consumirlo todo —declaró Akame con total solemnidad y los ojos rojos brillando como gemas, arrastrando a Pulchra hacia las bandejas humeantes sin perder un solo segundo.",
        "* * *",
        "En los jardines exteriores de roca, Rukia Kuchiki contemplaba embelesada una hilera de estanques menores decorados con grabados de conejos de porcelana blanca.",
        "—Qué tallado tan sublime... miren las orejas de estos pequeños espíritus del bosque —decía Rukia conmovida, admirando los adornos con una devoción casi sagrada.",
        "—¡Son simples conejos de adorno, Rukia, no te quedes pegada ahí! —le gritó Renji desde el puente de madera, enfundado en un yukata holgado con sus tatuajes rojizos al aire—. ¡Vamos a meternos al agua antes de que Ken se lance de bomba y la vacíe entera!",
        "—¡Ten modales, Renji! ¡La nobleza de este balneario exige compostura! —le recriminó ella dándole un golpe seco en la espinilla con el dobladillo de su toalla, antes de no poder evitar una risita y apresurarse hacia el estanque.",
        "Caesar King y Lucy se encargaron de organizar las alas de la suite con su habitual don de mando: el ala norte con los estanques termales para las chicas y el ala sur para los chicos, con una gran terraza común en el centro para el banquete nocturno.",
        "* * *",
        "El baño en las aguas termales fue una bendición.",
        "Sumergirse en el agua humeante y caliente hasta los hombros hizo que cada uno de los noventa kilos de músculo de mi cuerpo se relajara profundamente. La rigidez de los golpes del calabozo y la tensión de los exámenes se disolvieron en el vapor mineral.",
        "Lighter se recostó contra el borde de piedra con los ojos cerrados, sin sus gafas por una vez, luciendo una expresión de paz absoluta; Ken flotaba bocarriba echando vaho por la boca; y Natsuki y Yami discutían acaloradamente sobre si meterse a la piscina fría después del vapor era señal de valentía o de estupidez.",
        "Al caer la noche, todos nos reunimos en la terraza techada con batas cómodas de lino para compartir el festín de comida imperial.",
        "Kaoruko se sentó a mi lado con el cabello negro húmedo recogido en una coleta informal y las mejillas arreboladas por el calor del baño. Llevaba una bandeja pequeña con tazas de té dulce y brochetas glaseadas.",
        "—¿Descansaste bien, Seier-kun? —me preguntó en voz baja, ofreciéndome una taza con esa calidez que siempre me desarmaba.",
        "—Hacía años que no me sentía tan ligero —respondí, mirándola fijamente y apreciando lo hermosa y tranquila que se veía bajo la luz de las linternas de papel—. Valió la pena cada página que me hiciste memorizar.",
        "Kaoruko sonrió, y por un segundo apoyó su mano pequeña y tibia sobre mi brazo musculoso, dejándola allí con una confianza dulce y silenciosa.",
        "Alrededor nuestro, las risas de Burnice, los festines de Akame, los bostezos de Piper y las fanfarronadas de Ken llenaban la noche de una vida desbordante.",
        "Éramos la Clase F-2. Los que entraron por la puerta trasera... pero que disfrutaban de la cima con la frente en alto."
    ]
}

def compilar_capitulo_17():
    from compilar_novela import PROLOGO, CAPITULO_1
    from agregar_capitulo_2 import CAPITULO_2
    from agregar_capitulo_3 import CAPITULO_3
    from agregar_capitulo_4 import CAPITULO_4
    from agregar_capitulo_5 import CAPITULO_5
    from agregar_capitulo_6 import CAPITULO_6
    from agregar_capitulo_7 import CAPITULO_7
    from agregar_capitulo_8 import CAPITULO_8
    from agregar_capitulo_9 import CAPITULO_9
    from agregar_capitulo_10 import CAPITULO_10
    from agregar_capitulo_11 import CAPITULO_11
    from agregar_capitulo_12 import CAPITULO_12
    from agregar_capitulo_13 import CAPITULO_13
    from agregar_capitulo_14 import CAPITULO_14
    from agregar_capitulo_15 import CAPITULO_15
    from agregar_capitulo_16 import CAPITULO_16
    
    titulo = "El camino del reencarnado"
    subtitulo = "Crónicas de Sangre, Acero y Deseo"
    autor = "Seier Silva"
    capitulos = [PROLOGO, CAPITULO_1, CAPITULO_2, CAPITULO_3, CAPITULO_4, CAPITULO_5, CAPITULO_6, CAPITULO_7, CAPITULO_8, CAPITULO_9, CAPITULO_10, CAPITULO_11, CAPITULO_12, CAPITULO_13, CAPITULO_14, CAPITULO_15, CAPITULO_16, CAPITULO_17]
    
    # 1. Guardar texto plano del Capítulo 17
    txt_path = os.path.join(NOVELA_DIR, "capitulos", "17_Capitulo_17.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(f"=== {CAPITULO_17['header'].upper()}: {CAPITULO_17['title'].upper()} ===\n\n")
        f.write("\n\n".join(CAPITULO_17['paragraphs']))
    print(f"[OK] Texto plano guardado en: {txt_path}")
    
    # 2. Guardar Word individual del Capítulo 17
    docx_cap17 = os.path.join(NOVELA_DIR, "capitulos", "17_Capitulo_17.docx")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    r_head = p_title.add_run(f"{CAPITULO_17['header'].upper()}\n")
    r_head.font.name = "Georgia"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(140, 80, 160)
    
    r_tit = p_title.add_run(CAPITULO_17['title'])
    r_tit.font.name = "Georgia"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = RGBColor(30, 25, 45)

    for p_text in CAPITULO_17['paragraphs']:
        p_text = p_text.strip()
        if not p_text:
            continue
        if p_text in ["* * *", "✦   ✦   ✦"]:
            p_sep = doc.add_paragraph()
            p_sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sep.paragraph_format.space_before = Pt(10)
            p_sep.paragraph_format.space_after = Pt(10)
            r_s = p_sep.add_run("✦   ✦   ✦")
            r_s.font.name = "Georgia"
            r_s.font.size = Pt(11)
            r_s.font.color.rgb = RGBColor(160, 130, 180)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(5)
            if p_text.startswith("—") or p_text.startswith("- ") or p_text.startswith("«"):
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.first_line_indent = Inches(0.28)
            r = p.add_run(p_text)
            r.font.name = "Georgia"
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(28, 28, 32)
            
    doc.save(docx_cap17)
    print(f"[OK] Word individual Capítulo 17 guardado en: {docx_cap17}")
    
    # 3. Guardar Libro Maestro Completo (Prólogo + Cap 1 a 17)
    master_docx = os.path.join(NOVELA_DIR, "compilado", "El_Camino_del_Reencarnado.docx")
    create_novel_document(titulo, subtitulo, autor, capitulos, master_docx)
    print(f"[OK] Libro Maestro actualizado en: {master_docx}")

if __name__ == "__main__":
    compilar_capitulo_17()
